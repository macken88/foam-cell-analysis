"""作業データを表で整えて同時確定する画面。"""

from __future__ import annotations

from collections import Counter

from PySide6.QtCore import (
    QByteArray,
    QEvent,
    QItemSelection,
    QItemSelectionModel,
    Qt,
    QTimer,
    Signal,
)
from PySide6.QtGui import (
    QAction,
    QActionGroup,
    QColor,
    QKeySequence,
    QPalette,
    QShortcut,
    QStandardItem,
    QStandardItemModel,
    QUndoCommand,
    QUndoStack,
)
from PySide6.QtWidgets import (
    QAbstractItemView,
    QComboBox,
    QDialog,
    QFileDialog,
    QFrame,
    QGridLayout,
    QHBoxLayout,
    QHeaderView,
    QInputDialog,
    QLabel,
    QLineEdit,
    QMenu,
    QMessageBox,
    QPushButton,
    QSizePolicy,
    QSplitter,
    QStyle,
    QStyledItemDelegate,
    QStyleOptionViewItem,
    QTableView,
    QVBoxLayout,
    QWidget,
)

from ....services.models import DataItem, ImportCandidate
from ...context import DEFAULT_CHANNEL, AppContext
from ...settings import app_settings
from ...theme import Color, body_font, numeric_font, set_style
from ...widgets.image_convert import DisplayMode, array_to_pixmap, render
from ...widgets.image_view import ImageView
from ...widgets.marks import (
    USAGE_MARKS,
    CountChip,
    DisplayToggle,
    TagDelegate,
    display_mode_label,
)
from ...widgets.page_base import BasePage
from ...widgets.table import (
    add_row_context_menu,
    bind_button_action,
    fit_table_columns,
    restore_row_selection,
    setup_table,
)
from .dialogs import (
    AutoTriageDialog,
    ContinuousTriageDialog,
    DatasetFinalizeDialog,
    ImportDialog,
    ImportSettingsDialog,
)
from .finalize_thumbnails import DatasetVersionThumbnailWindow
from .table_model import (
    HEADERS,
    USAGE_TEXT,
    DataPreparationTableModel,
    DatasetHistoryModel,
    ValueComboDelegate,
)


class _EditCommand(QUndoCommand):
    """選択項目への一括変更を一つの取り消し操作にする。"""

    def __init__(self, page: DataPreparationPage, item_ids: list[str], changes: dict, text: str):
        super().__init__(text)
        self.page = page
        self.item_ids = list(item_ids)
        self.changes = changes
        self.before = {
            item.item_id: {key: getattr(item, key) for key in changes}
            for item in page.items
            if item.item_id in item_ids
        }
        self.before_flags = {
            item.item_id: (item.change, item.previous_change)
            for item in page.items
            if item.item_id in item_ids
        }

    def redo(self) -> None:
        self.page._apply_changes(self.item_ids, self.changes)

    def undo(self) -> None:
        for item_id, values in self.before.items():
            self.page.ctx.backend.update_item("all", item_id, **values)
        self.page._restore_change_flags(self.before_flags)
        self.page.refresh(
            self.item_ids,
            filter_membership_changed=bool(set(self.changes) & {"usage", "classification"}),
        )


class _UsageDelegate(TagDelegate):
    """用途タグを表示し、用途コンボで編集する。"""

    def createEditor(self, parent, option, index):
        editor = QComboBox(parent)
        editor.addItems(list(self.mapping))
        return editor

    def paint(self, painter, option, index) -> None:
        item = index.data(Qt.ItemDataRole.UserRole)
        if isinstance(item, DataItem) and item.change:
            color = (
                Color.SELECTION_CHANGED
                if option.state & QStyle.StateFlag.State_Selected
                else Color.CHANGED
            )
            painter.fillRect(option.rect, color)
            prepared = QStyleOptionViewItem(option)
            if option.state & QStyle.StateFlag.State_Selected:
                prepared.palette.setColor(QPalette.ColorRole.Highlight, QColor(color))
                prepared.palette.setColor(
                    QPalette.ColorRole.HighlightedText, QColor(Color.GRAPHITE)
                )
        else:
            prepared = option
        super().paint(painter, prepared, index)

    def setEditorData(self, editor, index) -> None:
        editor.setCurrentText(str(index.data(Qt.ItemDataRole.EditRole) or ""))

    def setModelData(self, editor, model, index) -> None:
        model.setData(index, editor.currentText(), Qt.ItemDataRole.EditRole)


class _ChangedRowDelegate(QStyledItemDelegate):
    """変更行の背景色を表のすべてのセルへ伝える。"""

    def paint(self, painter, option, index) -> None:
        item = index.data(Qt.ItemDataRole.UserRole)
        prepared = QStyleOptionViewItem(option)
        marker_color = index.data(Qt.ItemDataRole.ForegroundRole)
        if index.column() == 0 and marker_color is not None:
            prepared.palette.setColor(QPalette.ColorRole.Text, marker_color)
            prepared.palette.setColor(QPalette.ColorRole.HighlightedText, marker_color)
        if isinstance(item, DataItem) and item.change:
            color = (
                Color.SELECTION_CHANGED
                if option.state & QStyle.StateFlag.State_Selected
                else Color.CHANGED
            )
            # QSS 適用下では backgroundBrush が無視されるため、先に直接塗る
            painter.fillRect(option.rect, QColor(color))
            if option.state & QStyle.StateFlag.State_Selected:
                prepared.palette.setColor(QPalette.ColorRole.Highlight, QColor(color))
                prepared.palette.setColor(
                    QPalette.ColorRole.HighlightedText, QColor(Color.GRAPHITE)
                )
        super().paint(painter, prepared, index)


class _TriageCommand(QUndoCommand):
    """実行済みの自動振り分けを一括で取り消す。"""

    def __init__(
        self,
        page,
        before: dict[str, str],
        after: dict[str, str],
        before_flags: dict[str, tuple[str | None, str | None]],
    ):
        super().__init__("自動振り分け")
        self.page = page
        self.before = before
        self.after = after
        self.before_flags = before_flags
        self.first = True

    def redo(self):
        if self.first:
            self.first = False
            return
        self._apply(self.after)

    def undo(self):
        self._apply(self.before)
        self.page._restore_change_flags(self.before_flags)
        self.page.refresh(list(self.before), filter_membership_changed=True)

    def _apply(self, values):
        for item_id, usage in values.items():
            self.page.ctx.backend.update_item("all", item_id, usage=usage)
        self.page.refresh(list(values), filter_membership_changed=True)


class _ClassificationFilter(QComboBox):
    """件数付きの複数選択分類フィルター。"""

    changed = Signal()

    def __init__(self, parent=None) -> None:
        super().__init__(parent)
        self.setMinimumWidth(150)
        self.setFocusPolicy(Qt.FocusPolicy.StrongFocus)
        self.setEditable(True)
        self.lineEdit().setReadOnly(True)
        self.lineEdit().setFocusPolicy(Qt.FocusPolicy.NoFocus)
        self._items = QStandardItemModel(self)
        self.setModel(self._items)
        self._items.itemChanged.connect(self._item_changed)
        self.setPlaceholderText("画像分類")

    def set_classifications(
        self,
        values: list[tuple[str, int]],
        selected: set[str],
        total_count: int | None = None,
    ) -> None:
        self._items.blockSignals(True)
        self._items.clear()
        count = total_count if total_count is not None else sum(value for _, value in values)
        all_item = QStandardItem(f"すべて　{count}")
        all_item.setData(None, Qt.ItemDataRole.UserRole)
        all_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsUserCheckable)
        all_item.setCheckState(Qt.CheckState.Checked if not selected else Qt.CheckState.Unchecked)
        self._items.appendRow(all_item)
        for name, count in values:
            item = QStandardItem(f"{name}　{count}")
            item.setData(name, Qt.ItemDataRole.UserRole)
            item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsUserCheckable)
            item.setCheckState(
                Qt.CheckState.Checked if name in selected else Qt.CheckState.Unchecked
            )
            self._items.appendRow(item)
        self._items.blockSignals(False)
        self._sync_text()

    def selected_values(self) -> set[str]:
        return {
            str(item.data(Qt.ItemDataRole.UserRole))
            for row in range(1, self._items.rowCount())
            if (item := self._items.item(row)).checkState() == Qt.CheckState.Checked
        }

    def set_selected(self, selected: set[str]) -> None:
        self._items.blockSignals(True)
        self._items.item(0).setCheckState(
            Qt.CheckState.Checked if not selected else Qt.CheckState.Unchecked
        )
        for row in range(1, self._items.rowCount()):
            item = self._items.item(row)
            item.setCheckState(
                Qt.CheckState.Checked
                if item.data(Qt.ItemDataRole.UserRole) in selected
                else Qt.CheckState.Unchecked
            )
        self._items.blockSignals(False)
        self._sync_text()

    def _sync_text(self, *_args) -> None:
        selected = self.selected_values()
        label = "すべて" if not selected else "、".join(sorted(selected))
        self.setCurrentText(label)

    def showPopup(self) -> None:  # noqa: N802
        super().showPopup()

    def _item_changed(self, changed_item) -> None:
        if changed_item.row() == 0:
            if changed_item.checkState() == Qt.CheckState.Checked:
                self._items.blockSignals(True)
                for row in range(1, self._items.rowCount()):
                    self._items.item(row).setCheckState(Qt.CheckState.Unchecked)
                self._items.blockSignals(False)
            elif not self.selected_values():
                self._items.blockSignals(True)
                changed_item.setCheckState(Qt.CheckState.Checked)
                self._items.blockSignals(False)
        elif changed_item.checkState() == Qt.CheckState.Checked:
            all_item = self._items.item(0)
            if all_item.checkState() == Qt.CheckState.Checked:
                self._items.blockSignals(True)
                all_item.setCheckState(Qt.CheckState.Unchecked)
                self._items.blockSignals(False)
        elif not self.selected_values():
            self._items.blockSignals(True)
            self._items.item(0).setCheckState(Qt.CheckState.Checked)
            self._items.blockSignals(False)
        self._sync_text()
        self.changed.emit()


class DataPreparationPage(BasePage):
    """作業表、プレビュー、絞り込み、キー操作をまとめて表示する。"""

    def __init__(self, ctx: AppContext, parent=None, show_heading: bool = True) -> None:
        super().__init__(
            ctx, "作業中データ", "画像を確認して用途・分類・品質を決めます。", parent, show_heading
        )
        self.shortcuts = ctx.shortcuts
        self.undo_stack = QUndoStack(self)
        self.items = ctx.backend.get_working_items()
        self.model = DataPreparationTableModel(self)
        self._selected_ids: list[str] = []
        self._syncing_selection = False
        self._selection_invalidated = False
        self._current_item_id: str | None = None
        root = self.content_layout
        root.setSpacing(6)
        self.base_label = QLabel()
        self.base_label.setFont(numeric_font(9))
        self.import_button = QPushButton("取り込み")
        self.auto_button = QPushButton("自動振り分け")
        self.excel_button = QPushButton("Excel ▾")
        self.other_button = QPushButton("その他 ▾")
        self.finalize_button = QPushButton("確定")
        self.finalize_button.setProperty("primary", True)
        self.import_button.hide()
        self.auto_button.hide()
        self.excel_button.hide()
        self.other_button.hide()
        self.finalize_error_button = QPushButton()
        self.finalize_error_button.setVisible(False)
        self.finalize_error_button.clicked.connect(lambda: self._set_error_filter(True))
        self.tab_tools = QWidget()
        self.tab_tools_layout = QHBoxLayout(self.tab_tools)
        self.tab_tools_layout.setContentsMargins(0, 0, 0, 0)
        self.tab_tools_layout.addWidget(self.base_label)
        self.tab_tools_layout.addWidget(self.finalize_error_button)
        self.tab_tools_layout.addWidget(self.finalize_button)
        usage_row = QHBoxLayout()
        usage_row.setSpacing(0)
        self.chips: dict[str, CountChip] = {}
        chip_defs = [
            ("all", "すべて", "plain"),
            ("unassigned", "未振り分け", "unassigned"),
            ("train", "学習", "train"),
            ("val", "検証", "val"),
            ("excluded", "不採用", "excluded"),
        ]
        for key, label, usage in chip_defs:
            chip = CountChip(label, 0, usage)
            chip.setFocusPolicy(Qt.FocusPolicy.NoFocus)
            chip.clicked.connect(lambda checked=False, name=key: self._filter_usage(name))
            usage_row.addWidget(chip)
            self.chips[key] = chip
        self.chips["all"].setChecked(True)
        root.addLayout(self._labeled_filter_row("用途", usage_row))

        filter_row = QHBoxLayout()
        filter_row.setSpacing(8)
        self.error_filter = CountChip("⚠ エラー", 0, "error")
        self.error_filter.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        self.error_filter.toggled.connect(self._set_error_filter)
        self.chips["errors"] = self.error_filter
        self.changed_filter = QPushButton("変更あり 0")
        self.changed_filter.setCheckable(True)
        self.changed_filter.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        set_style(self.changed_filter, role="filterToggle", usage="changed")
        self.changed_filter.toggled.connect(self._set_changed_filter)
        filter_row.addWidget(self.error_filter)
        filter_row.addWidget(self.changed_filter)
        filter_row.addSpacing(12)
        self.class_filter = _ClassificationFilter()
        self.class_filter.changed.connect(self._classification_filter_changed)
        filter_row.addWidget(QLabel("画像分類"))
        filter_row.addWidget(self.class_filter)
        self.source_combo = QComboBox()
        self.source_combo.setMinimumWidth(155)
        self.source_combo.addItem("すべて", None)
        self.source_combo.currentIndexChanged.connect(self._source_changed)
        filter_row.addWidget(QLabel("取り込み元"))
        filter_row.addWidget(self.source_combo)
        self.search = QLineEdit()
        self.search.setMinimumWidth(175)
        self.search.setPlaceholderText("識別子・ファイル名を検索")
        self.search.textChanged.connect(self._search_changed)
        filter_row.addWidget(self.search, 1)
        root.addLayout(self._labeled_filter_row("絞り込み", filter_row, stretch_controls=True))
        for chip in self.chips.values():
            chip.toggled.connect(lambda checked, target=chip: self._style_chip(target, checked))
            self._style_chip(chip, chip.isChecked())
        for name, key in (
            ("all", "filter_all"),
            ("unassigned", "filter_unassigned"),
            ("train", "filter_train"),
            ("val", "filter_val"),
            ("excluded", "filter_excluded"),
        ):
            self.chips[name].setToolTip(self.shortcuts.display_key(self.shortcuts[key]))
        self.splitter = QSplitter(Qt.Orientation.Horizontal)
        self.table = QTableView()
        self.table.setModel(self.model)
        self.table.setItemDelegate(_ChangedRowDelegate(self.table))
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        self.table.verticalHeader().hide()
        self.table.setSortingEnabled(False)
        self.table.horizontalHeader().setStretchLastSection(False)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        setup_table(
            self.table,
            selection_mode=QAbstractItemView.SelectionMode.ExtendedSelection,
        )
        usage_mapping = {
            USAGE_TEXT[key]: (values[1], values[2], values[3])
            for key, values in USAGE_MARKS.items()
        }
        self.table.setItemDelegateForColumn(2, _UsageDelegate(usage_mapping, self.table))
        self.table.setItemDelegateForColumn(
            3, ValueComboDelegate(["未設定", *ctx.backend.classifications], self.table)
        )
        self.table.setItemDelegateForColumn(
            4, ValueComboDelegate(["未設定", "良", "可", "不良"], self.table)
        )
        self.model.edit_requested.connect(self._table_edit_requested)
        table_panel = QWidget()
        table_layout = QVBoxLayout(table_panel)
        table_layout.setContentsMargins(0, 0, 0, 0)
        table_head = QHBoxLayout()
        self.visible_count = QLabel("0 件中 0 件を表示")
        self.column_button = QPushButton("表示する列 ▾", table_panel)
        self.column_button.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        self.column_menu = QMenu(self.column_button)
        self.column_button.setMenu(self.column_menu)
        table_head.addWidget(self.visible_count)
        table_head.addStretch(1)
        table_head.addWidget(self.column_button)
        table_layout.addLayout(table_head)
        table_layout.addWidget(self.table, 1)
        self.splitter.addWidget(table_panel)
        self.preview_panel = QWidget()
        preview = QVBoxLayout(self.preview_panel)
        preview.setContentsMargins(8, 0, 0, 0)
        self.preview_meta = QLabel("画像を選択してください")
        self.preview_meta.setFont(body_font(13))
        preview.addWidget(self.preview_meta)
        self.preview_details = QLabel("")
        self.preview_details.setFont(numeric_font(9))
        set_style(self.preview_details, role="note")
        preview.addWidget(self.preview_details)
        self.image_view = ImageView()
        self.image_view.setMinimumWidth(155)
        self.display_toggle = DisplayToggle(ctx.display, label_scope="ground_truth")
        self.display_toggle.alternate_selected.connect(self._show_preview)
        preview.addWidget(self.display_toggle)
        preview.addWidget(self.image_view, 1)
        self.image_caption = QLabel("画像表示形式")
        set_style(self.image_caption, role="note")
        preview.addWidget(self.image_caption)
        edits_box = QFrame()
        edits_box.setProperty("role", "filterPanel")
        edits_layout = QVBoxLayout(edits_box)
        edits_layout.setContentsMargins(10, 8, 10, 8)
        self.selection_note = QLabel("選択中の 0 件を変更")
        edits_layout.addWidget(self.selection_note)
        edits = QGridLayout()
        edits.setHorizontalSpacing(8)
        edits.setVerticalSpacing(4)
        self._edit_grid_layout = edits
        self._edit_pair_layouts = []
        self.usage_combo = QComboBox()
        for key, label in USAGE_TEXT.items():
            self.usage_combo.addItem(label, key)
        self.class_combo = QComboBox()
        self.class_combo.addItem("未設定", None)
        for value in ctx.backend.classifications:
            self.class_combo.addItem(value, value)
        self.quality_combo = QComboBox()
        self.quality_combo.addItem("未設定", None)
        for value in ("良", "可", "不良"):
            self.quality_combo.addItem(value, value)
        self.mask_combo = QComboBox()
        for combo in (self.usage_combo, self.class_combo, self.quality_combo, self.mask_combo):
            combo.setSizeAdjustPolicy(QComboBox.SizeAdjustPolicy.AdjustToContents)
            combo.setSizePolicy(QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Fixed)
        edits.setSpacing(20)
        for combo in (self.usage_combo, self.class_combo, self.quality_combo, self.mask_combo):
            combo.currentIndexChanged.connect(
                lambda _index, control=combo: self._edit_selection(control)
            )
        self.edit_combos = {}
        for position, (label, combo) in enumerate(
            (
                ("用途", self.usage_combo),
                ("分類", self.class_combo),
                ("品質", self.quality_combo),
                ("正解ラベル版", self.mask_combo),
            )
        ):
            pair = QHBoxLayout()
            pair.setSpacing(6)
            pair.addWidget(QLabel(label))
            pair.addWidget(combo)
            self._edit_pair_layouts.append(pair)
            edits.addLayout(pair, 0, position)
            self.edit_combos[label] = combo
        edits_layout.addLayout(edits)
        preview.addWidget(edits_box)
        self.splitter.addWidget(self.preview_panel)
        self.splitter.setStretchFactor(0, 3)
        self.splitter.setStretchFactor(1, 2)
        self.preview_panel.setMinimumWidth(220)
        self.splitter.setSizes([1050, 438])
        root.addWidget(self.splitter, 1)
        self._make_menus()
        self.table.selectionModel().selectionChanged.connect(
            lambda *_: self._selection_changed(self.table)
        )
        self.table.selectionModel().currentChanged.connect(self._current_changed)
        self.table.installEventFilter(self)
        self.table.viewport().installEventFilter(self)
        self._install_filter_shortcuts()
        self.image_view.installEventFilter(self)
        self.image_view.viewport().installEventFilter(self)
        self.preview_panel.installEventFilter(self)
        self._layout_preview_edits()
        bind_button_action(self.import_button, self.menu_action_map["import"])
        bind_button_action(self.auto_button, self.menu_action_map["auto_triage"])
        bind_button_action(self.finalize_button, self.menu_action_map["finalize"])
        settings = app_settings()
        saved_widths = settings.value("dataPreparation/columnWidths")
        self._has_saved_column_widths = False
        self._initial_column_widths_done = False
        if saved_widths:
            self._has_saved_column_widths = self.table.horizontalHeader().restoreState(
                QByteArray.fromBase64(saved_widths.encode())
            )
            for column in range(self.model.columnCount()):
                if self.table.columnWidth(column) < 24:
                    self.table.setColumnWidth(column, 24)
        for column, key in enumerate(
            (
                "state",
                "filename",
                "usage",
                "classification",
                "quality",
                "identifier",
                "source",
                "mask",
            )
        ):
            self.table.setColumnHidden(column, not self.column_action_map[key].isChecked())
        self._saving_column_widths = True
        self.table.horizontalHeader().sectionResized.connect(self._save_column_widths)
        self.refresh()
        QTimer.singleShot(0, self._initialize_column_widths)
        self.shortcuts.changed.connect(self._shortcuts_changed)
        self.ctx.display.changed.connect(self._display_preference_changed)
        self._shortcuts_changed()

    def _initialize_column_widths(self) -> None:
        """初回表示で内容に合わせて列幅を決める。

        余白を特定の列へ足すと、ウィンドウが後から狭くなったときに
        ほかの列が画面外へ押し出されるため、内容幅だけで決める。
        """
        if self._initial_column_widths_done or not self.isVisible():
            return
        self._initial_column_widths_done = True
        if not self._has_saved_column_widths:
            header = self.table.horizontalHeader()
            self._saving_column_widths = False
            fit_table_columns(self.table)
            header.resizeSection(0, 34)
            header.resizeSection(1, max(header.sectionSizeHint(1), header.sectionSize(1) - 6))
            header.resizeSection(3, max(header.sectionSize(3), 96))
            self._saving_column_widths = True

    @staticmethod
    def _labeled_filter_row(
        label: str, controls: QHBoxLayout, *, stretch_controls: bool = False
    ) -> QHBoxLayout:
        row = QHBoxLayout()
        row.setSpacing(8)
        title = QLabel(label)
        title.setFixedWidth(58)
        row.addWidget(title)
        row.addLayout(controls, int(stretch_controls))
        if not stretch_controls:
            row.setAlignment(controls, Qt.AlignmentFlag.AlignLeft)
        return row

    def _make_menus(self) -> None:
        excel = QMenu(self)
        excel.addAction("Excel 出力", self.export_excel)
        excel.addAction("Excel 取込", self.import_excel)
        self.excel_button.setMenu(excel)
        other = QMenu(self)
        other.addAction("連続振り分け", self.open_triage)
        other.addAction("キー割り当て一覧", self.open_keymap)
        display_menu = other.addMenu("原画像と切り替える表示")
        self.display_actions = {}
        for name in sorted(self.ctx.display.MODES):
            action = display_menu.addAction(display_mode_label(name, "ground_truth"))
            action.setCheckable(True)
            action.setChecked(name == self.ctx.display.value)
            action.triggered.connect(
                lambda checked=False, value=name: self.ctx.display.set_value(value)
            )
            self.display_actions[name] = action
        other.addAction("元に戻す", self.undo_stack.undo)
        other.addAction("やり直す", self.undo_stack.redo)
        self.mask_revision_action = QAction("正解ラベル版を追加", self)
        other.addAction(self.mask_revision_action)
        self.mask_revision_action.triggered.connect(self.import_mask_revision)
        self.archive_action = QAction("アーカイブを作成", self)
        other.addAction(self.archive_action)
        self.archive_action.triggered.connect(self.create_archive)
        self.other_button.setMenu(other)
        self._shortcut_actions = {}
        for action, callback in (
            ("import", self.import_data),
            ("auto_triage", self.auto_triage),
            ("export_excel", self.export_excel),
            ("import_excel", self.import_excel),
            ("finalize", self.finalize),
            ("search", self.search.setFocus),
        ):
            shortcut = QAction(self)
            shortcut.setShortcut(QKeySequence(self.shortcuts[action]))
            shortcut.setShortcutContext(Qt.ShortcutContext.WidgetWithChildrenShortcut)
            shortcut.triggered.connect(callback)
            self.addAction(shortcut)
            self._shortcut_actions[action] = shortcut
        for name in (
            "filter_all",
            "filter_unassigned",
            "filter_train",
            "filter_val",
            "filter_excluded",
            "filter_errors",
        ):
            action = QAction(self)
            action.setShortcut(QKeySequence(self.shortcuts[name]))
            self._shortcut_actions[name] = action
        self._scoped_shortcuts = []
        for name, callback, targets in (
            ("undo", self.undo_stack.undo, (self.table, self.image_view)),
            ("redo", self.undo_stack.redo, (self.table, self.image_view)),
            ("select_all", self.table.selectAll, (self.table,)),
            ("clear_selection", self._clear_selection_or_filter, (self.table,)),
        ):
            action = QAction(self)
            action.triggered.connect(callback)
            self._shortcut_actions[name] = action
            for target in targets:
                shortcut = QShortcut(QKeySequence(self.shortcuts[name]), target)
                shortcut.setContext(Qt.ShortcutContext.WidgetWithChildrenShortcut)
                shortcut.activated.connect(callback)
                self._scoped_shortcuts.append(shortcut)
        labels = {
            "import": "画像を取り込む",
            "auto_triage": "自動振り分け",
            "export_excel": "Excel に出力",
            "import_excel": "Excel から取り込む",
            "finalize": "データセットを確定",
            "undo": "元に戻す",
            "redo": "やり直す",
            "search": "検索",
            "select_all": "すべて選択",
            "clear_selection": "選択を解除",
        }
        self.menu_action_map = {}
        for key, label in labels.items():
            shortcut_action = self._shortcut_actions.get(key)
            if shortcut_action is not None:
                action = QAction(self)
                action.triggered.connect(shortcut_action.trigger)
            else:
                action = QAction(self)
                action.triggered.connect(
                    (lambda: self.table.selectAll())
                    if key == "select_all"
                    else self._clear_selection_or_filter
                )
            action.setText(self._menu_text(label, key if key in self.shortcuts.mapping else None))
            self.menu_action_map[key] = action
        self.menu_action_map["mask_revision"] = self.mask_revision_action
        self.menu_action_map["archive"] = self.archive_action
        self.menu_action_map["triage"] = QAction("連続振り分け\tEnter", self)
        self.menu_action_map["triage"].triggered.connect(self.open_triage)
        self.menu_action_map["error_filter"] = QAction("エラーのある画像を表示", self)
        self.menu_action_map["error_filter"].triggered.connect(lambda: self._set_error_filter(True))
        self.menu_action_map["keymap"] = QAction("キー割り当て", self)
        self.menu_action_map["keymap"].triggered.connect(self.open_keymap)
        self.usage_menu = QMenu("用途を変更", self)
        self.usage_action_map = {}
        for key, label, value in (
            ("usage_unassigned", "未振り分け", "unassigned"),
            ("usage_train", "学習", "train"),
            ("usage_val", "検証", "val"),
            ("usage_excluded", "不採用", "excluded"),
        ):
            action = QAction(self._menu_text(label, key), self)
            action.triggered.connect(
                lambda _checked=False, v=value: self._change(self._selected_ids, usage=v)
            )
            self.usage_menu.addAction(action)
            self.usage_action_map[key] = action
        self.class_menu = QMenu("画像分類を変更", self)
        self.class_action_map = {}
        self._classification_actions = []
        self._static_class_actions = {}
        for key, label, value in (
            ("class_clear", "未設定", None),
            ("class_dialog", "一覧から選ぶ", "dialog"),
        ):
            action = QAction(self._menu_text(label, key), self)
            action.triggered.connect(
                lambda _checked=False, v=value: (
                    self.choose_classification()
                    if v == "dialog"
                    else self._change(self._selected_ids, classification=v)
                )
            )
            self._static_class_actions[key] = action
        self._class_menu_separator = self.class_menu.addSeparator()
        for action in self._static_class_actions.values():
            self.class_menu.addAction(action)
        self._refresh_classification_menu()
        self.quality_menu = QMenu("品質を変更", self)
        self.quality_action_map = {}
        for key, label, value in (
            ("quality_good", "良", "good"),
            ("quality_ok", "可", "acceptable"),
            ("quality_bad", "不良", "bad"),
        ):
            action = QAction(self._menu_text(label, key), self)
            action.triggered.connect(
                lambda _checked=False, v=value: self._change(self._selected_ids, quality=v)
            )
            self.quality_menu.addAction(action)
            self.quality_action_map[key] = action
        self.filter_menu = QMenu("絞り込み", self)
        self.filter_action_map = {}
        for key, label in (
            ("filter_all", "すべて"),
            ("filter_unassigned", "未振り分け"),
            ("filter_train", "学習"),
            ("filter_val", "検証"),
            ("filter_excluded", "不採用"),
        ):
            action = QAction(self._menu_text(label, key), self)
            action.setCheckable(True)
            action.triggered.connect(
                lambda _checked=False, k=key: self._filter_usage(k.removeprefix("filter_"))
            )
            self.filter_menu.addAction(action)
            self.filter_action_map[key] = action
        self.usage_filter_group = QActionGroup(self)
        self.usage_filter_group.setExclusive(True)
        for action in self.filter_action_map.values():
            self.usage_filter_group.addAction(action)
        self.filter_menu.addSeparator()
        for key, label in (("filter_errors", "⚠ エラー"), ("filter_changed", "変更あり")):
            action = QAction(label, self)
            action.setCheckable(True)
            action.toggled.connect(
                self._set_error_filter if key == "filter_errors" else self._set_changed_filter
            )
            self.filter_menu.addAction(action)
            self.filter_action_map[key] = action
        self.filter_menu.addSeparator()
        self.classification_filter_menu = self.filter_menu.addMenu("画像分類")
        self.classification_filter_actions = {}
        self.filter_all_classes_action = self.classification_filter_menu.addAction("すべて")
        self.filter_all_classes_action.setCheckable(True)
        self.filter_all_classes_action.triggered.connect(lambda: self._set_classifications(set()))
        for classification in self.ctx.backend.classifications:
            action = self.classification_filter_menu.addAction(classification)
            action.setCheckable(True)
            action.toggled.connect(
                lambda checked, value=classification: self._toggle_class_filter(value, checked)
            )
            self.classification_filter_actions[classification] = action
        self._make_column_actions()
        self.display_menu = QMenu("原画像と切り替える表示", self)
        for action in self.display_actions.values():
            self.display_menu.addAction(action)
        self.image_actions = {}
        for key, label, callback in (
            ("display_mode", "原画像と切り替える", self._cycle_preview_mode),
            ("previous_image", "前の画像", lambda: self._move_image(-1)),
            ("next_image", "次の画像", lambda: self._move_image(1)),
            ("zoom_in", "拡大", lambda: self.image_view.zoom_by(1.2)),
            ("zoom_out", "縮小", lambda: self.image_view.zoom_by(1 / 1.2)),
            ("fit_view", "全体表示", self.image_view.fit_image),
        ):
            action = QAction(f"{label}\t{self.shortcuts.display_key(self.shortcuts[key])}", self)
            action.triggered.connect(callback)
            self.image_actions[key] = action
        self.context_menu = QMenu(self)
        for action in (
            self.usage_menu.menuAction(),
            self.class_menu.menuAction(),
            self.quality_menu.menuAction(),
        ):
            self.context_menu.addAction(action)
        self.context_menu.addSeparator()
        for action in (
            self.menu_action_map["triage"],
            self.menu_action_map["undo"],
            self.menu_action_map["redo"],
        ):
            self.context_menu.addAction(action)
        add_row_context_menu(self.table, self.context_menu)

    def _make_column_actions(self) -> None:
        settings = app_settings()
        saved = settings.value("dataPreparation/visibleColumns")
        visible = (
            set(saved) if saved else {"state", "filename", "usage", "classification", "quality"}
        )
        self.column_action_map = {}
        self.column_menu.clear()
        for column, key in enumerate(
            (
                "state",
                "filename",
                "usage",
                "classification",
                "quality",
                "identifier",
                "source",
                "mask",
            )
        ):
            action = self.column_menu.addAction(HEADERS[column])
            action.setCheckable(True)
            action.setChecked(key in visible)
            action.toggled.connect(
                lambda checked, col=column, name=key: self._set_column_visible(col, name, checked)
            )
            self.column_action_map[key] = action
            self.table.setColumnHidden(column, key not in visible)

    def _refresh_filter_classification_menu(self) -> None:
        self.classification_filter_menu.clear()
        self.classification_filter_actions = {}
        self.filter_all_classes_action = self.classification_filter_menu.addAction("すべて")
        self.filter_all_classes_action.setCheckable(True)
        self.filter_all_classes_action.triggered.connect(lambda: self._set_classifications(set()))
        for classification in self.ctx.backend.classifications:
            action = self.classification_filter_menu.addAction(classification)
            action.setCheckable(True)
            action.toggled.connect(
                lambda checked, value=classification: self._toggle_class_filter(value, checked)
            )
            self.classification_filter_actions[classification] = action

    def _set_column_visible(self, column: int, key: str, visible: bool) -> None:
        self.table.setColumnHidden(column, not visible)
        names = [
            "state",
            "filename",
            "usage",
            "classification",
            "quality",
            "identifier",
            "source",
            "mask",
        ]
        current = {name for name, action in self.column_action_map.items() if action.isChecked()}
        app_settings().setValue(
            "dataPreparation/visibleColumns", [name for name in names if name in current]
        )
        self._save_column_widths()

    def _menu_text(self, label: str, shortcut: str | None) -> str:
        if shortcut:
            return f"{label}\t{self.shortcuts.display_key(self.shortcuts[shortcut])}"
        return label

    def menu_actions(self):
        return {
            "file": [
                self.menu_action_map["import"],
                self.menu_action_map["mask_revision"],
                None,
                self.menu_action_map["export_excel"],
                self.menu_action_map["import_excel"],
                None,
                self.menu_action_map["archive"],
            ],
            "edit": [
                self.menu_action_map["undo"],
                self.menu_action_map["redo"],
                None,
                self.menu_action_map["select_all"],
                self.menu_action_map["clear_selection"],
                self.menu_action_map["search"],
                None,
                self.usage_menu.menuAction(),
                self.class_menu.menuAction(),
                self.quality_menu.menuAction(),
            ],
            "view": [
                None,
                self.column_menu.menuAction(),
                self.filter_menu.menuAction(),
                None,
                self.image_actions["display_mode"],
                self.image_actions["previous_image"],
                self.image_actions["next_image"],
                None,
                self.image_actions["zoom_in"],
                self.image_actions["zoom_out"],
                self.image_actions["fit_view"],
            ],
            "dataset": [
                self.menu_action_map["auto_triage"],
                self.menu_action_map["triage"],
                None,
                self.menu_action_map["error_filter"],
                self.menu_action_map["finalize"],
            ],
            "tools": [],
        }

    def _refresh_classification_menu(self) -> None:
        for action in self._classification_actions:
            self.class_menu.removeAction(action)
        self._classification_actions = []
        self.class_action_map = {}
        classifications = list(self.ctx.backend.classifications)
        for number, classification in enumerate(classifications, 1):
            key = f"class_{number}"
            label = self._menu_text(classification, key) if number <= 9 else classification
            action = QAction(label, self)
            action.triggered.connect(
                lambda _checked=False, value=classification: self._change(
                    self._selected_ids, classification=value
                )
            )
            self.class_menu.insertAction(self._class_menu_separator, action)
            self._classification_actions.append(action)
            self.class_action_map[key] = action
        for key, action in self._static_class_actions.items():
            self.class_action_map[key] = action

    def tab_tools_widget(self):
        return self.tab_tools

    def _add_key_action(self, name: str, callback) -> QAction:
        action = QAction(self)
        action.setShortcut(QKeySequence(self.shortcuts[name]))
        action.setShortcutContext(Qt.ShortcutContext.WidgetWithChildrenShortcut)
        action.triggered.connect(callback)
        self.addAction(action)
        return action

    def _install_filter_shortcuts(self) -> None:
        """件数チップのキー操作を表にフォーカスがある間有効にする。"""
        self._filter_shortcuts = {}
        for name in (
            "filter_all",
            "filter_unassigned",
            "filter_train",
            "filter_val",
            "filter_excluded",
            "filter_errors",
        ):
            shortcut = QShortcut(QKeySequence(self.shortcuts[name]), self.table)
            shortcut.setContext(Qt.ShortcutContext.WidgetWithChildrenShortcut)
            if name == "filter_errors":
                shortcut.activated.connect(lambda: self.error_filter.toggle())
            else:
                shortcut.activated.connect(
                    lambda target=name: self._filter_usage(target.removeprefix("filter_"))
                )
            self._filter_shortcuts[name] = shortcut

    def _shortcuts_changed(self) -> None:
        """共有キー変更を画面の操作・表示へ反映する。"""
        for name, action in self._shortcut_actions.items():
            action.setShortcut(
                QKeySequence(self.shortcuts[name])
                if name not in {"undo", "redo", "select_all", "clear_selection"}
                else QKeySequence()
            )
        for shortcut, name in zip(
            self._scoped_shortcuts,
            ("undo", "undo", "redo", "redo", "select_all", "clear_selection"),
            strict=True,
        ):
            shortcut.setKey(QKeySequence(self.shortcuts[name]))
        for name in self.shortcuts.mapping:
            if name in self.menu_action_map:
                label = self.menu_action_map[name].text().split("\t", 1)[0]
                self.menu_action_map[name].setText(self._menu_text(label, name))
            for action_map in (
                self.usage_action_map,
                self.class_action_map,
                self.quality_action_map,
                self.filter_action_map,
                self.image_actions,
            ):
                if name in action_map:
                    label = action_map[name].text().split("\t", 1)[0]
                    action_map[name].setText(self._menu_text(label, name))
        for name, shortcut in self._filter_shortcuts.items():
            shortcut.setKey(QKeySequence(self.shortcuts[name]))
        for chip, name in (
            ("all", "filter_all"),
            ("unassigned", "filter_unassigned"),
            ("train", "filter_train"),
            ("val", "filter_val"),
            ("excluded", "filter_excluded"),
        ):
            self.chips[chip].setToolTip(self.shortcuts.display_key(self.shortcuts[name]))
        for button, name in (
            (self.import_button, "import"),
            (self.auto_button, "auto_triage"),
            (self.finalize_button, "finalize"),
        ):
            button.setToolTip(
                f"{button.text().strip()}（{self.shortcuts.display_key(self.shortcuts[name])}）"
            )
        self.search.setToolTip(f"検索（{self.shortcuts.display_key(self.shortcuts['search'])}）")
        self.excel_button.setToolTip(
            "Excel 入出力（"
            f"{self.shortcuts.display_key(self.shortcuts['export_excel'])} / "
            f"{self.shortcuts.display_key(self.shortcuts['import_excel'])}）"
        )

    def _apply_changes(self, item_ids: list[str], changes: dict) -> None:
        if "usage" in changes and len(changes) == 1:
            self.ctx.backend.set_usage(item_ids, changes["usage"])
        else:
            self.ctx.backend.update_items(item_ids, **changes)
        self.refresh(
            item_ids,
            filter_membership_changed=bool(set(changes) & {"usage", "classification"}),
        )
        if item_ids:
            names = {
                "train": "学習",
                "val": "検証",
                "excluded": "不採用",
                "unassigned": "未振り分け",
            }
            label = names.get(
                changes.get("usage"),
                str(changes.get("classification", changes.get("quality", "変更"))),
            )
            first = next((item.item_id for item in self.items if item.item_id in item_ids), "")
            suffix = f" ほか {len(item_ids) - 1}件" if len(item_ids) > 1 else ""
            self.ctx.status.show_message(
                f"{first}{suffix} を {label} にしました（"
                f"{self.shortcuts.display_key(self.shortcuts['undo'])} で元に戻す）"
            )

    def _restore_change_flags(self, values: dict[str, tuple[str | None, str | None]]) -> None:
        """Undo 実行前の差分表示状態へ戻す。"""
        by_id = {item.item_id: item for item in self.ctx.backend.get_working_items()}
        for item_id, (change, previous_change) in values.items():
            if item_id in by_id:
                by_id[item_id].change = change
                by_id[item_id].previous_change = previous_change

    def _table_edit_requested(self, item_id: str, field: str, value) -> None:
        self._change([item_id], **{field: value})

    def _change(self, item_ids: list[str], **changes) -> None:
        if item_ids:
            self.undo_stack.push(_EditCommand(self, item_ids, changes, "作業データの変更"))

    def _selection_changed(self, source=None) -> None:
        if self._syncing_selection:
            return
        self._syncing_selection = True
        ids = [
            item.item_id
            for index in self.table.selectionModel().selectedRows()
            if (item := self.model.item_at(index.row())) is not None
        ]
        if ids:
            self._selection_invalidated = False
        elif self._selected_ids:
            self._selection_invalidated = True
        self._selected_ids = ids
        self.selection_note.setText(f"選択中の {len(ids)} 件を変更")
        self.set_menu_action_enabled(self.mask_revision_action, bool(ids))
        self.mask_revision_action.setToolTip("作業中データの行を選ぶと使えます" if not ids else "")
        for action in (
            *self.usage_action_map.values(),
            *self.class_action_map.values(),
            *self.quality_action_map.values(),
        ):
            self.set_menu_action_enabled(action, bool(ids))
            action.setToolTip("作業中データの行を選ぶと使えます" if not ids else "")
        self._show_preview()
        self._syncing_selection = False

    def _current_changed(self, current, _previous) -> None:
        if self._syncing_selection or not current.isValid():
            return
        item = self.model.item_at(current.row())
        self._current_item_id = item.item_id if item else None

    def _edit_selection(self, control) -> None:
        if not self._selected_ids or control.currentIndex() < 0:
            return
        if control is self.usage_combo:
            self._change(self._selected_ids, usage=control.currentData())
        elif control is self.class_combo:
            self._change(self._selected_ids, classification=control.currentData())
        elif control is self.quality_combo:
            self._change(self._selected_ids, quality=control.currentData())
        elif control is self.mask_combo:
            self._change(self._selected_ids, selected_mask_revision=control.currentText())

    def _show_preview(self, *_args) -> None:
        if not self._selected_ids:
            self.image_view.set_image(None)
            self._preview_signature = None
            self.preview_meta.setText("画像を選択してください")
            self.preview_details.setText("")
            self.image_caption.setText("画像表示形式")
            self.mask_combo.clear()
            for control in (
                self.usage_combo,
                self.class_combo,
                self.quality_combo,
                self.mask_combo,
            ):
                control.setEnabled(False)
            return
        item = next((item for item in self.items if item.item_id == self._selected_ids[0]), None)
        if not item:
            return
        self.preview_meta.setText(item.source_filename)
        self.preview_details.setText(
            f"{item.item_id}　取り込み元 {item.source_folder}　正解ラベル版 "
            f"{item.selected_mask_revision or 'なし'}"
        )
        self.usage_combo.blockSignals(True)
        self.class_combo.blockSignals(True)
        self.quality_combo.blockSignals(True)
        self.mask_combo.blockSignals(True)
        self.usage_combo.setEnabled(True)
        self.class_combo.setEnabled(True)
        self.quality_combo.setEnabled(True)
        self.usage_combo.setCurrentIndex(self.usage_combo.findData(item.usage))
        self.class_combo.setCurrentIndex(max(0, self.class_combo.findData(item.classification)))
        self.quality_combo.setCurrentIndex(max(0, self.quality_combo.findData(item.quality)))
        selected_items = [
            candidate for candidate in self.items if candidate.item_id in self._selected_ids
        ]
        common_revisions = set(item.mask_revisions)
        for selected_item in selected_items[1:]:
            common_revisions.intersection_update(selected_item.mask_revisions)
        revisions = [revision for revision in item.mask_revisions if revision in common_revisions]
        self.mask_combo.clear()
        self.mask_combo.addItems(revisions)
        self.mask_combo.setEnabled(bool(revisions))
        if item.selected_mask_revision in revisions:
            self.mask_combo.setCurrentText(item.selected_mask_revision)
        self.usage_combo.blockSignals(False)
        self.class_combo.blockSignals(False)
        self.quality_combo.blockSignals(False)
        self.mask_combo.blockSignals(False)
        channel = DEFAULT_CHANNEL
        image = self.ctx.backend.get_item_image("all", item.item_id, channel)
        mode = self._preview_mode()
        labels = (
            self.ctx.backend.get_item_mask("all", item.item_id, item.selected_mask_revision)
            if item.mask_revisions
            else None
        )
        signature = (tuple(self._selected_ids), item.selected_mask_revision, channel, mode)
        skip_unchanged = getattr(self, "_preserve_preview_on_refresh", False)
        if not skip_unchanged or signature != getattr(self, "_preview_signature", None):
            self.image_view.set_image(array_to_pixmap(render(image, labels, mode)))
            self._preview_signature = signature
        self.image_caption.setText(
            "原画像"
            if mode == DisplayMode.IMAGE
            else display_mode_label(self.ctx.display.value, "ground_truth")
        )

    def _filter_usage(self, name: str) -> None:
        mapping = {
            "all": None,
            "unassigned": {"unassigned"},
            "train": {"train"},
            "val": {"val"},
            "excluded": {"excluded"},
        }
        self.model.usages = mapping[name]
        selected_before = self._selected_ids[:]
        current_before = self._current_item_id
        self.refresh_views()
        self._selected_ids = selected_before
        self._sync_selection_to_visible(current_before)
        for key in ("all", "unassigned", "train", "val", "excluded"):
            self.chips[key].setChecked(key == name)
        self._sync_filter_menu()

    def _set_error_filter(self, enabled: bool) -> None:
        self.model.errors_only = enabled
        self.error_filter.setChecked(enabled)
        selected_before = self._selected_ids[:]
        current_before = self._current_item_id
        self.refresh_views()
        self._selected_ids = selected_before
        self._sync_selection_to_visible(current_before)
        self._sync_filter_menu()

    def _set_changed_filter(self, enabled: bool) -> None:
        self.model.changed_only = enabled
        self.changed_filter.setChecked(enabled)
        selected_before = self._selected_ids[:]
        current_before = self._current_item_id
        self.refresh_views()
        self._selected_ids = selected_before
        self._sync_selection_to_visible(current_before)
        self._sync_filter_menu()

    def _classification_filter_changed(self) -> None:
        self._set_classifications(self.class_filter.selected_values())

    def _toggle_class_filter(self, classification: str, checked: bool) -> None:
        selected = set(self.model.classifications)
        if checked:
            selected.add(classification)
        else:
            selected.discard(classification)
        self._set_classifications(selected)

    def _set_classifications(self, selected: set[str]) -> None:
        self.model.classifications = set(selected)
        self.class_filter.set_selected(selected)
        for name, action in self.classification_filter_actions.items():
            action.blockSignals(True)
            action.setChecked(name in selected)
            action.blockSignals(False)
        self.filter_all_classes_action.blockSignals(True)
        self.filter_all_classes_action.setChecked(not selected)
        self.filter_all_classes_action.blockSignals(False)
        selected_before = self._selected_ids[:]
        current_before = self._current_item_id
        selected_before = self._selected_ids[:]
        self.refresh_views()
        self._selected_ids = selected_before
        self._sync_selection_to_visible(current_before)
        self._sync_filter_menu()

    def _sync_filter_menu(self) -> None:
        for key, name in (
            ("filter_all", "all"),
            ("filter_unassigned", "unassigned"),
            ("filter_train", "train"),
            ("filter_val", "val"),
            ("filter_excluded", "excluded"),
        ):
            self.filter_action_map[key].blockSignals(True)
            usage = None if name == "all" else {name}
            self.filter_action_map[key].setChecked(self.model.usages == usage)
            self.filter_action_map[key].blockSignals(False)
        for key, enabled in (
            ("filter_errors", self.model.errors_only),
            ("filter_changed", self.model.changed_only),
        ):
            action = self.filter_action_map[key]
            action.blockSignals(True)
            action.setChecked(enabled)
            action.blockSignals(False)
        for name, action in self.classification_filter_actions.items():
            action.blockSignals(True)
            action.setChecked(name in self.model.classifications)
            action.blockSignals(False)
        self.filter_all_classes_action.blockSignals(True)
        self.filter_all_classes_action.setChecked(not self.model.classifications)
        self.filter_all_classes_action.blockSignals(False)

    @staticmethod
    def _style_chip(chip: CountChip, checked: bool) -> None:
        """テーマの名前付き状態でチップの文字色を切り替える。"""
        state = "activeChipText" if checked else ""
        set_style(chip.name_label, state=state)
        set_style(chip.count_label, state=state)

    def _filter_class(self, value: str) -> None:
        self._toggle_class_filter(value, value not in self.model.classifications)

    def _source_changed(self, index: int) -> None:
        self.model.source_folder = self.source_combo.currentData()
        selected_before = self._selected_ids[:]
        current_before = self._current_item_id
        self.refresh_views()
        self._selected_ids = selected_before
        self._sync_selection_to_visible(current_before)

    def _search_changed(self, text: str) -> None:
        self.model.query = text
        selected_before = self._selected_ids[:]
        current_before = self._current_item_id
        self.refresh_views()
        self._selected_ids = selected_before
        self._sync_selection_to_visible(current_before)

    def _sync_selection_to_visible(self, current_id: str | None = None) -> None:
        """絞り込み後に選択とプレビューを表示行だけへそろえる。"""
        visible = self.model.visible_items()
        visible_ids = {item.item_id for item in visible}
        scroll_value = self.table.verticalScrollBar().value()
        if current_id is not None:
            self._current_item_id = current_id
        previous_ids = self._selected_ids[:]
        self._selected_ids = [item_id for item_id in previous_ids if item_id in visible_ids]
        if previous_ids and not self._selected_ids:
            self._selection_invalidated = True
        elif self._selected_ids:
            self._selection_invalidated = False
        if (
            not self._selected_ids
            and visible
            and not self._selection_invalidated
            and not previous_ids
        ):
            self._selected_ids = [visible[0].item_id]
        self._restore_table_selection()
        self.table.verticalScrollBar().setValue(scroll_value)
        self._show_preview()

    def _restore_table_selection(self) -> None:
        """_selected_ids のすべての行を表で選び直す（複数選択を保つ）。

        selectRow は拡張選択でも既存の選択を消すため、更新のたびに
        選択が 1 件へ縮んでいた。QItemSelection でまとめて選ぶ。
        """
        visible = self.model.visible_items()
        rows = {item.item_id: row for row, item in enumerate(visible)}
        selection = QItemSelection()
        for item_id in self._selected_ids:
            row = rows.get(item_id)
            if row is not None:
                selection.select(
                    self.model.index(row, 0),
                    self.model.index(row, self.model.columnCount() - 1),
                )
        model = self.table.selectionModel()
        self._syncing_selection = True
        try:
            model.select(selection, QItemSelectionModel.SelectionFlag.ClearAndSelect)
            selected_ids = set(self._selected_ids)
            current_id = self._current_item_id
            current_row = rows.get(current_id) if current_id in selected_ids else None
            if current_row is None and self._selected_ids:
                current_id = self._selected_ids[0]
                current_row = rows.get(current_id)
            if current_row is not None:
                model.setCurrentIndex(
                    self.model.index(
                        current_row,
                        model.currentIndex().column() if model.currentIndex().isValid() else 0,
                    ),
                    QItemSelectionModel.SelectionFlag.NoUpdate,
                )
                self._current_item_id = current_id
            elif not self._selected_ids:
                model.clearCurrentIndex()
                self._current_item_id = None
        finally:
            self._syncing_selection = False
        self.selection_note.setText(f"選択中の {len(self._selected_ids)} 件を変更")

    def _select_only_row(self, row: int, item_id: str) -> None:
        """拡張選択状態でも対象行だけを選ぶ。"""
        self._syncing_selection = True
        self.table.clearSelection()
        self.table.selectRow(row)
        self._selected_ids = [item_id]
        self._selection_invalidated = False
        self._syncing_selection = False

    def refresh_views(self) -> None:
        self.model.filters_changed()
        self.visible_count.setText(
            f"{len(self.items)} 件中 {len(self.model.visible_items())} 件を表示"
        )

    def refresh(
        self, item_ids: list[str] | None = None, *, filter_membership_changed: bool = False
    ) -> None:
        scroll_value = self.table.verticalScrollBar().value()
        previous_ids = self._selected_ids[:]
        self._refresh_classification_menu()
        self._refresh_filter_classification_menu()
        self.items = self.ctx.backend.get_working_items()
        report = self.ctx.backend.validate_items()
        errors = {issue.item_id: issue.message for issue in report.errors}
        if item_ids is None:
            self.model.set_items(self.items, errors)
        else:
            self.model.update_item_errors(
                errors,
                set(item_ids),
                refresh_layout=(
                    filter_membership_changed or self.model.errors_only or self.model.changed_only
                ),
            )
        self.base_label.setText(
            f"ベース版　学習 {self.ctx.backend.working['all'].base_train_version} / "
            f"検証 {self.ctx.backend.working['all'].base_val_version}"
        )
        counts = Counter(item.usage for item in self.items)
        for key, value in (
            ("all", len(self.items)),
            ("unassigned", counts["unassigned"]),
            ("train", counts["train"]),
            ("val", counts["val"]),
            ("excluded", counts["excluded"]),
            ("errors", len(errors)),
            ("changed", sum(item.change is not None for item in self.items)),
        ):
            if key in self.chips:
                self.chips[key].count_label.setText(str(value))
        self.error_filter.count_label.setText(str(len(errors)))
        self.changed_filter.setText(
            f"変更あり {sum(item.change is not None for item in self.items)}"
        )
        class_counts = Counter(item.classification for item in self.items)
        self.class_filter.set_classifications(
            [(name, class_counts[name]) for name in self.ctx.backend.classifications],
            self.model.classifications,
            len(self.items),
        )
        selected_folder = self.source_combo.currentData()
        folders = sorted({item.source_folder for item in self.items})
        self.source_combo.blockSignals(True)
        self.source_combo.clear()
        self.source_combo.addItem("すべて", None)
        for folder in folders:
            self.source_combo.addItem(folder, folder)
        self.source_combo.setCurrentIndex(max(0, self.source_combo.findData(selected_folder)))
        self.source_combo.blockSignals(False)
        self.visible_count.setText(
            f"{len(self.items)} 件中 {len(self.model.visible_items())} 件を表示"
        )
        self.finalize_button.setEnabled(not errors)
        self.set_menu_action_enabled(self.menu_action_map["finalize"], not errors)
        finalize_tip = f"確定（{self.shortcuts.display_key(self.shortcuts['finalize'])}）"
        if errors:
            finalize_tip += "　整合性エラーを解消してください"
        self.finalize_button.setToolTip(finalize_tip)
        self.menu_action_map["finalize"].setToolTip(
            f"エラー {len(errors)} 件を直すと確定できます" if errors else ""
        )
        self.finalize_error_button.setText(f"⚠ エラー {len(errors)} 件を直すと確定できます")
        set_style(self.finalize_error_button, usage="error")
        self.finalize_error_button.setVisible(bool(errors))
        self._sync_filter_menu()
        visible_ids = {item.item_id for item in self.model.visible_items()}
        self._selected_ids = [item_id for item_id in self._selected_ids if item_id in visible_ids]
        if previous_ids and not self._selected_ids:
            self._selection_invalidated = True
        elif self._selected_ids:
            self._selection_invalidated = False
        if (
            not self._selected_ids
            and self.model.visible_items()
            and not self._selection_invalidated
            and not previous_ids
        ):
            self._selected_ids = [self.model.visible_items()[0].item_id]
        self._restore_table_selection()
        self._show_preview()
        self.table.verticalScrollBar().setValue(scroll_value)
        self.set_menu_action_enabled(self.mask_revision_action, bool(self._selected_ids))
        self.set_menu_action_enabled(
            self.archive_action, bool(self.ctx.backend.list_dataset_versions())
        )
        self.archive_action.setToolTip(
            "データセット版を確定するとアーカイブを作成できます"
            if not self.ctx.backend.list_dataset_versions()
            else ""
        )

    def import_data(self) -> None:
        dialog = ImportDialog(self)
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return
        try:
            candidates = self.ctx.backend.scan_import_folders(
                dialog.image_dirs, dialog.mask_dir, None
            )
        except ValueError as error:
            QMessageBox.warning(self, "取り込みエラー", str(error))
            return
        folders: dict[str, list[ImportCandidate]] = {}
        for item in candidates:
            folders.setdefault(item.source_relpath.rsplit("/", 1)[0], []).append(item)
        settings = ImportSettingsDialog(self, folders)
        if settings.exec() != QDialog.DialogCode.Accepted:
            return
        imported_items = self.ctx.backend.import_folders(candidates, settings.values)
        self.model.usages = {item.usage for item in imported_items}
        self.model.source_folder = None
        self.source_combo.setCurrentIndex(0)
        self.refresh()

    def auto_triage(self) -> None:
        dialog = AutoTriageDialog(
            self, self.items, selected_ids=self._selected_ids, backend=self.ctx.backend
        )
        if dialog.exec() == QDialog.DialogCode.Accepted:
            settings = dialog.values()
            target_selected = settings.pop("target_selected")
            target_ids = self._selected_ids if target_selected else None
            before = {
                item.item_id: item.usage
                for item in self.items
                if (target_ids is not None or item.usage == "unassigned")
                and (target_ids is None or item.item_id in target_ids)
            }
            before_flags = {
                item.item_id: (item.change, item.previous_change)
                for item in self.items
                if item.item_id in before
            }
            result = self.ctx.backend.apply_auto_split(settings, target_ids)
            after = {item.item_id: item.usage for item in self.items if item.item_id in before}
            self.undo_stack.push(_TriageCommand(self, before, after, before_flags))
            self._change_status(f"{sum(result.values())} 件を振り分けました")

    def _change_status(self, message: str) -> None:
        self.ctx.status.show_message(message)
        self.refresh()

    def open_triage(self) -> None:
        visible_before = self.model.visible_items()
        selected_ids = [
            visible_before[index.row()].item_id
            for index in self.table.selectionModel().selectedRows()
            if index.row() < len(visible_before)
        ]
        targets = self.model.visible_items()
        dialog = ContinuousTriageDialog(
            self,
            targets,
            lambda item_id, **changes: self._change([item_id], **changes),
            self.ctx.backend,
            self.shortcuts,
        )
        dialog.exec()
        current_id = dialog.current_item_id
        self.refresh()
        visible = self.model.visible_items()
        visible_rows = {item.item_id: row for row, item in enumerate(visible)}
        self._selected_ids = [item_id for item_id in selected_ids if item_id in visible_rows]
        self._restore_table_selection()
        if current_id and current_id in visible_rows:
            self.table.selectionModel().setCurrentIndex(
                self.model.index(visible_rows[current_id], 0),
                QItemSelectionModel.SelectionFlag.NoUpdate,
            )
            self._show_preview()

    def finalize(self) -> None:
        dialog = DatasetFinalizeDialog(self, self.ctx.backend)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            versions = dialog.created
            label = (
                " と ".join(item.version for item in versions) if versions else "変更はありません"
            )
            self._change_status(f"{label} を確定しました")

    def export_excel(self) -> None:
        path, _ = QFileDialog.getSaveFileName(
            self, "作業データを Excel に出力", "working_dataset.xlsx", "Excel (*.xlsx)"
        )
        if not path:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.worksheet.datavalidation import DataValidation

            book = Workbook()
            sheet = book.active
            sheet.title = "作業データ"
            sheet.append(
                [
                    "データ識別子",
                    "取り込み元フォルダ",
                    "元ファイル名",
                    "用途",
                    "画像分類",
                    "品質",
                    "正解ラベル版",
                ]
            )
            for item in self.model.visible_items():
                sheet.append(
                    [
                        item.item_id,
                        item.source_folder,
                        item.source_filename,
                        USAGE_TEXT[item.usage],
                        item.classification or "",
                        item.quality or "",
                        item.selected_mask_revision,
                    ]
                )
            for col, formula in (
                (4, '"未振り分け,学習,検証,不採用"'),
                (5, '"分類A,分類B,分類C"'),
                (6, '"良,可,不良"'),
            ):
                rule = DataValidation(type="list", formula1=formula, allow_blank=True)
                sheet.add_data_validation(rule)
                if sheet.max_row >= 2:
                    rule.add(f"{chr(64 + col)}2:{chr(64 + col)}{sheet.max_row}")
            book.save(path)
        except ImportError:
            QMessageBox.warning(self, "Excel 出力", "Excel 出力に必要なライブラリがありません")

    def import_excel(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self, "作業データを Excel から取り込む", "", "Excel (*.xlsx)"
        )
        if not path:
            return
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet = workbook.active
            rows = list(sheet.iter_rows(min_row=2, values_only=True))
            updates = []
            by_id = {item.item_id: item for item in self.items}
            usage_map = {v: k for k, v in USAGE_TEXT.items()}
            for row in rows:
                if row[0] not in by_id:
                    continue
                if row[3] not in usage_map:
                    raise ValueError(f"{row[0]} の用途が正しくありません")
                updates.append(
                    (
                        str(row[0]),
                        {
                            "usage": usage_map[row[3]],
                            "classification": row[4] or None,
                            "quality": row[5] or None,
                            "selected_mask_revision": row[6] or "",
                        },
                    )
                )
            command = _ExcelCommand(self, updates)
            self.undo_stack.push(command)
        except (ImportError, OSError, ValueError) as error:
            QMessageBox.warning(self, "Excel 取込エラー", str(error))
        finally:
            if "workbook" in locals():
                workbook.close()

    def open_keymap(self) -> None:
        from ...keymap_dialog import show_keymap_window

        show_keymap_window(self, self.ctx)

    def choose_classification(self) -> None:
        """選択項目へ分類をまとめて設定する。"""
        if not self._selected_ids:
            return
        value, accepted = QInputDialog.getItem(
            self, "画像分類", "画像分類", self.ctx.backend.classifications, 0, False
        )
        if accepted:
            self._change(self._selected_ids, classification=value)

    def show_key_help(self) -> None:
        """現在のキー割り当てを表示する。"""
        self.open_keymap()

    def _clear_selection_or_filter(self) -> None:
        """最初は選択を解除し、次はすべての絞り込みを解除する。"""
        if self._selected_ids:
            self.table.clearSelection()
            self._selected_ids = []
            self._show_preview()
        else:
            self.search.clear()
            self.source_combo.setCurrentIndex(0)
            self._filter_usage("all")
            self._set_error_filter(False)
            self._set_changed_filter(False)
            self._set_classifications(set())

    def import_mask_revision(self) -> None:
        """選択中の項目へ新しい正解ラベル版を追加する。"""
        if not self._selected_ids:
            return
        created = []
        for item_id in self._selected_ids:
            created.append(self.ctx.backend.add_mask_revision("all", item_id))
        self.refresh(self._selected_ids)
        self.ctx.status.show_message(f"{len(created)} 件に正解ラベル版を追加しました")

    def create_archive(self) -> None:
        """最新の学習用・検証用版のアーカイブを作成する。"""
        versions = self.ctx.backend.list_dataset_versions()
        latest_by_purpose = {}
        for version in versions:
            latest_by_purpose[version.purpose] = version
        if not latest_by_purpose:
            QMessageBox.information(self, "アーカイブ作成", "確定済みデータセット版がありません")
            return
        output_path = QFileDialog.getExistingDirectory(self, "アーカイブの保存先を選択")
        if not output_path:
            return
        for version in latest_by_purpose.values():
            self.ctx.backend.record_archive_result(version.version, output_path)
        self.ctx.status.show_message("データセットのアーカイブを作成しました")

    def _cycle_preview_mode(self) -> None:
        self.display_toggle.set_alternate(not self.display_toggle.is_alternate)
        self.ctx.status.show_message(
            "表示形式: "
            + (
                display_mode_label(self.ctx.display.value, "ground_truth")
                if self.display_toggle.is_alternate
                else "原画像"
            )
        )

    def _preview_mode(self) -> DisplayMode:
        """二択表示の選択を描画モードへ変換する。"""
        if not self.display_toggle.is_alternate:
            return DisplayMode.IMAGE
        return {
            "オーバーレイ": DisplayMode.OVERLAY,
            "インスタンスラベル": DisplayMode.INSTANCE_LABEL,
            "二値マスク": DisplayMode.BINARY,
        }[self.ctx.display.value]

    def _move_image(self, delta: int) -> None:
        """作業表で前後の画像行を選ぶ。"""
        visible = self.model.visible_items()
        if not visible:
            return
        current = next(
            (index for index, item in enumerate(visible) if item.item_id in self._selected_ids), 0
        )
        row = max(0, min(len(visible) - 1, current + delta))
        self._select_only_row(row, visible[row].item_id)
        self._show_preview()

    def _display_preference_changed(self, name: str) -> None:
        """共有表示名と全プレビューを更新する。"""
        for key, action in self.display_actions.items():
            action.setChecked(key == name)
        self._show_preview()

    def _save_column_widths(self, *_args) -> None:
        """作業表の利用者設定幅を保存する。"""
        if getattr(self, "_saving_column_widths", False):
            app_settings().setValue(
                "dataPreparation/columnWidths",
                self.table.horizontalHeader().saveState().toBase64().data().decode(),
            )

    def eventFilter(self, watched, event) -> bool:
        if watched is self.preview_panel and event.type() == QEvent.Type.Resize:
            self._layout_preview_edits(event.size().width())
            return super().eventFilter(watched, event)
        if event.type() == QEvent.Type.KeyPress and watched in (
            self.table,
            self.table.viewport(),
            self.image_view,
            self.image_view.viewport(),
        ):
            from PySide6.QtWidgets import QComboBox, QLineEdit

            focus = self.focusWidget()
            if isinstance(focus, (QComboBox, QLineEdit)):
                return super().eventFilter(watched, event)
            if self.shortcuts.matches("previous_image", event):
                self._move_image(-1)
                return True
            if self.shortcuts.matches("next_image", event):
                self._move_image(1)
                return True
            if self.shortcuts.matches("zoom_in", event):
                self.image_view.zoom_by(1.2)
                return True
            if self.shortcuts.matches("zoom_out", event):
                self.image_view.zoom_by(1 / 1.2)
                return True
            if self.shortcuts.matches("fit_view", event):
                self.image_view.fit_image()
                return True
            for action, field, value in (
                ("usage_train", "usage", "train"),
                ("usage_val", "usage", "val"),
                ("usage_excluded", "usage", "excluded"),
                ("usage_unassigned", "usage", "unassigned"),
                ("quality_good", "quality", "良"),
                ("quality_ok", "quality", "可"),
                ("quality_bad", "quality", "不良"),
            ):
                if self.shortcuts.matches(action, event) and self._selected_ids:
                    self._change(self._selected_ids, **{field: value})
                    return True
            if self.shortcuts.matches("display_mode", event):
                self._cycle_preview_mode()
                return True
            if self.shortcuts.matches("class_dialog", event) and self._selected_ids:
                self.choose_classification()
                return True
            if self.shortcuts.matches("help", event) or event.key() == Qt.Key.Key_F1:
                self.show_key_help()
                return True
            if self.shortcuts.matches("previous_unassigned", event):
                self.select_next_unassigned(-1)
                return True
            if self.shortcuts.matches("next_unassigned", event):
                self.select_next_unassigned()
                return True
            if self.shortcuts.matches("triage_view", event):
                self.open_triage()
                return True
            if event.key() == Qt.Key.Key_Escape:
                self._clear_selection_or_filter()
                return True
            if self._selected_ids:
                if self.shortcuts.matches("class_clear", event):
                    self._change(self._selected_ids, classification=None)
                    return True
                for index in range(9):
                    if self.shortcuts.matches(f"class_{index + 1}", event):
                        if index < len(self.ctx.backend.classifications):
                            self._change(
                                self._selected_ids,
                                classification=self.ctx.backend.classifications[index],
                            )
                        return True
        return super().eventFilter(watched, event)

    def _layout_preview_edits(self, width: int | None = None) -> None:
        if not hasattr(self, "_edit_grid_layout"):
            return
        available = width if width is not None else self.preview_panel.width()
        per_row = 2 if available < 420 else 4
        for index, pair in enumerate(self._edit_pair_layouts):
            self._edit_grid_layout.removeItem(pair)
            self._edit_grid_layout.addLayout(pair, index // per_row, index % per_row)

    def select_next_unassigned(self, direction: int = 1) -> None:
        """次または前の未振り分け画像へ選択を移す。"""
        visible = self.model.visible_items()
        if not visible:
            return
        current = next(
            (index for index, item in enumerate(visible) if item.item_id in self._selected_ids), -1
        )
        for offset in range(1, len(visible) + 1):
            index = (current + direction * offset) % len(visible)
            if visible[index].usage == "unassigned":
                item = visible[index]
                self._select_only_row(index, item.item_id)
                self._show_preview()
                return

    def on_enter(self, params: dict) -> None:
        self.refresh()

    def refresh_on_activate(self) -> None:
        """編集中のセルを保ち、それ以外は最新データへ更新する。"""
        if self.table.state() == QAbstractItemView.State.EditingState:
            return
        self._preserve_preview_on_refresh = True
        try:
            self.refresh()
        finally:
            self._preserve_preview_on_refresh = False

    def refresh_menu_actions(self) -> None:
        """作業データの状態に応じたメニュー項目の有効状態を更新する。"""
        self.refresh_on_activate()


class _ExcelCommand(QUndoCommand):
    """Excel 取込の変更全体を取り消し可能にする。"""

    def __init__(self, page, updates):
        super().__init__("Excel からの変更")
        self.page = page
        self.updates = updates
        self.before = {
            item.item_id: (
                item.usage,
                item.classification,
                item.quality,
                item.selected_mask_revision,
            )
            for item in page.items
            if any(item.item_id == key for key, _ in updates)
        }
        self.before_flags = {
            item.item_id: (item.change, item.previous_change)
            for item in page.items
            if any(item.item_id == key for key, _ in updates)
        }

    def redo(self):
        for item_id, changes in self.updates:
            self.page.ctx.backend.update_item("all", item_id, **changes)
        self.page.refresh()

    def undo(self):
        for item_id, values in self.before.items():
            self.page.ctx.backend.update_item(
                "all",
                item_id,
                usage=values[0],
                classification=values[1],
                quality=values[2],
                selected_mask_revision=values[3],
            )
        self.page._restore_change_flags(self.before_flags)
        self.page.refresh()


class DatasetHistoryPage(BasePage):
    """学習用・検証用を一緒に確認する版履歴ページ。"""

    def __init__(self, ctx: AppContext, parent=None, show_heading: bool = False) -> None:
        super().__init__(
            ctx,
            "データセット版履歴",
            "確定済みの学習用版・検証用版を確認します。",
            parent,
            show_heading,
        )
        self.model = DatasetHistoryModel(self)
        self.thumbnail_windows = {}
        controls = QHBoxLayout()
        self.thumbnail_button = QPushButton("サムネイルで確認")
        self.thumbnail_action = QAction("選択した版をサムネイルで確認", self)
        self.thumbnail_action.triggered.connect(self.open_thumbnails)
        bind_button_action(self.thumbnail_button, self.thumbnail_action)
        controls.addStretch(1)
        controls.addWidget(self.thumbnail_button)
        self.content_layout.addLayout(controls)
        self.table = QTableView()
        self.table.setModel(self.model)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.table.verticalHeader().hide()
        self.table.setSortingEnabled(False)
        self.table.horizontalHeader().setStretchLastSection(False)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        setup_table(self.table)
        self.content_layout.addWidget(self.table)
        self.table.doubleClicked.connect(lambda _index: self.open_thumbnails())
        self.table.selectionModel().selectionChanged.connect(self._update_thumbnail_action)
        self.history_context_menu = QMenu(self)
        self.history_context_menu.addAction(self.thumbnail_action)
        add_row_context_menu(self.table, self.history_context_menu)
        self.refresh()

    def menu_actions(self):
        return {"dataset": [None, self.thumbnail_action]}

    def _update_thumbnail_action(self, *_args) -> None:
        enabled = bool(self.table.selectionModel().selectedRows())
        self.set_menu_action_enabled(self.thumbnail_action, enabled)
        self.thumbnail_action.setToolTip(
            "データセット版履歴で行を選ぶと使えます" if not enabled else ""
        )

    def refresh(self) -> None:
        self.model.set_versions(self.ctx.backend.list_dataset_versions())
        fit_table_columns(self.table)
        self._update_thumbnail_action()

    def on_enter(self, params: dict) -> None:
        self.refresh()

    def open_thumbnails(self) -> None:
        """選択した確定版のサムネイルウィンドウを開く。"""
        rows = self.table.selectionModel().selectedRows()
        if not rows:
            return
        version = self.model.versions[rows[0].row()]
        window = self.thumbnail_windows.get(version.version)
        if window is None:
            linked = next(
                (
                    item
                    for item in self.ctx.backend.list_dataset_versions("val")
                    if item.version == version.base_validation_version
                ),
                None,
            )
            window = DatasetVersionThumbnailWindow(self, self.ctx.backend, version, linked)
            self.thumbnail_windows[version.version] = window
        window.show()
        window.raise_()
        window.activateWindow()

    def refresh_on_activate(self) -> None:
        """版一覧を更新し、選択中の版を再選択する。"""
        selected_versions = [
            self.model.index(index.row(), 0).data()
            for index in self.table.selectionModel().selectedRows()
        ]
        current_version = (
            self.model.index(self.table.currentIndex().row(), 0).data()
            if self.table.currentIndex().isValid()
            else None
        )
        scroll_value = self.table.verticalScrollBar().value()
        self.refresh()
        version_rows = {version.version: row for row, version in enumerate(self.model.versions)}
        rows = {version_rows[version] for version in selected_versions if version in version_rows}
        restore_row_selection(self.table, rows, version_rows.get(current_version))
        self.table.verticalScrollBar().setValue(scroll_value)
