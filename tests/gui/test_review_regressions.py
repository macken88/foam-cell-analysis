"""GUI レビュー指摘の操作経路に対する回帰確認。"""

from unittest.mock import patch

import pytest
from PySide6.QtCore import QEvent, QPoint, Qt, QTimer
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QDialog, QDialogButtonBox, QFileDialog, QMessageBox

from foam_cell_analysis.gui.context import AppContext
from foam_cell_analysis.gui.jobs import JobManager
from foam_cell_analysis.gui.keymap_dialog import KeymapDialog, ShortcutKeyEdit
from foam_cell_analysis.gui.modes.comparison.candidates_page import CandidatesPage
from foam_cell_analysis.gui.modes.data_preparation.dialogs import (
    ContinuousTriageDialog,
    DatasetFinalizeDialog,
    ImportDialog,
    ImportSettingsDialog,
)
from foam_cell_analysis.gui.modes.inference.page import InferenceInput
from foam_cell_analysis.gui.navigation import ModeId, Navigator, PageId
from foam_cell_analysis.services.mock.backend import MockBackend


def _click_menu_action(qapp, button, action):
    """ボタンからメニューを開き、指定項目の表示位置をクリックする。"""
    QTest.mouseClick(button, Qt.MouseButton.LeftButton)
    qapp.processEvents()
    menu = button.menu()
    QTest.mouseClick(menu, Qt.MouseButton.LeftButton, pos=menu.actionGeometry(action).center())
    qapp.processEvents()


def _click_table_row(qapp, page, item_id, modifiers=Qt.KeyboardModifier.NoModifier):
    """作業表の識別子セルをクリックして行を選ぶ。"""
    row = next(
        index for index, item in enumerate(page.model.visible_items()) if item.item_id == item_id
    )
    cell = page.model.index(row, 1)
    QTest.mouseClick(
        page.table.viewport(),
        Qt.MouseButton.LeftButton,
        modifiers,
        page.table.visualRect(cell).center(),
    )
    qapp.processEvents()


def _choose_combo_option(qapp, combo, text):
    """コンボを開き、表示できる場合は項目をマウスで選ぶ。"""
    target = combo.findText(text)
    assert target >= 0
    combo.window().activateWindow()
    combo.setFocus()
    qapp.processEvents()
    QTest.keyClick(combo, Qt.Key.Key_F4)
    qapp.processEvents()
    if combo.view().isVisible():
        model_index = combo.model().index(target, 0)
        rect = combo.view().visualRect(model_index)
        QTest.mouseMove(combo.view().viewport(), rect.center())
        QTest.mouseClick(combo.view().viewport(), Qt.MouseButton.LeftButton, pos=rect.center())
    else:
        delta = target - combo.currentIndex()
        key = Qt.Key.Key_Down if delta > 0 else Qt.Key.Key_Up
        for _ in range(abs(delta)):
            QTest.keyClick(combo, key)
        QTest.keyClick(combo, Qt.Key.Key_Return)
    qapp.processEvents()
    assert combo.currentText() == text


def test_review_01_mask_revision_choices_are_common_and_backend_rejects_foreign_revision(
    shell, qapp
):
    page = shell.page(PageId.DATA_PREPARATION)
    rev2 = next(item for item in page.items if "rev_002" in item.mask_revisions)
    rev1_only = next(
        item
        for item in page.items
        if item.item_id != rev2.item_id and item.mask_revisions == ["rev_001"]
    )
    _click_table_row(qapp, page, rev2.item_id)
    _click_table_row(qapp, page, rev1_only.item_id, Qt.KeyboardModifier.ControlModifier)
    assert [page.mask_combo.itemText(i) for i in range(page.mask_combo.count())] == ["rev_001"]
    with pytest.raises(ValueError, match="マスク版"):
        shell.ctx.backend.update_item("all", rev1_only.item_id, selected_mask_revision="rev_002")


def test_review_02_filter_click_clears_selection_of_hidden_rows(shell, qapp):
    page = shell.page(PageId.DATA_PREPARATION)
    hidden = next(item for item in page.items if item.usage == "train")
    original_classification = hidden.classification
    _click_table_row(qapp, page, hidden.item_id)
    QTest.mouseClick(page.chips["val"], Qt.MouseButton.LeftButton)
    page.table.setFocus()
    QTest.keyClick(page.table, Qt.Key.Key_2)
    qapp.processEvents()
    assert hidden.classification == original_classification
    assert all(
        item_id in {item.item_id for item in page.model.visible_items()}
        for item_id in page._selected_ids
    )
    assert page._selected_ids[0] in page.preview_details.text()


def test_review_03_start_training_mouse_click_keeps_confirmation(shell, qapp):
    shell.navigate(PageId.TRAINING)
    page = shell.page(PageId.TRAINING)
    with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.No) as ask:
        QTest.mouseClick(page.start_button, Qt.MouseButton.LeftButton)
        qapp.processEvents()
    assert ask.call_count == 1
    assert shell.ctx.jobs.running_count == 0


def test_review_04_usage_and_quality_shortcuts_match_key_sequences(shell, qapp):
    page = shell.page(PageId.DATA_PREPARATION)
    item = next(item for item in page.items if item.usage == "unassigned")
    _click_table_row(qapp, page, item.item_id)
    page.table.setFocus()
    for key, usage in (
        (Qt.Key.Key_Q, "train"),
        (Qt.Key.Key_W, "val"),
        (Qt.Key.Key_E, "excluded"),
        (Qt.Key.Key_R, "unassigned"),
    ):
        QTest.keyClick(page.table, key)
        qapp.processEvents()
        assert item.usage == usage
    for key, quality in ((Qt.Key.Key_Z, "良"), (Qt.Key.Key_X, "可"), (Qt.Key.Key_C, "不良")):
        QTest.keyClick(page.table, key)
        qapp.processEvents()
        assert item.quality == quality


def test_review_05_inference_cannot_start_a_second_job_mid_run(shell, qapp):
    page = shell.page(PageId.INFERENCE)
    page.inputs = [InferenceInput(f"{index}.png", "分類A", "model_007") for index in range(3)]
    page.output_path = "C:/output"
    page._refresh_table()
    QTest.mouseClick(page.run_button, Qt.MouseButton.LeftButton)
    assert shell.ctx.jobs.running_count == 1
    assert not page.run_button.isEnabled()
    QTest.mouseClick(page.run_button, Qt.MouseButton.LeftButton)
    assert shell.ctx.jobs.running_count == 1


def test_review_06_candidate_page_refreshes_validation_versions_from_home(shell, qapp):
    shell.navigate(PageId.CANDIDATES)
    page = shell.page(PageId.CANDIDATES)
    for candidate in shell.ctx.backend.get_working_items():
        if candidate.usage in {"train", "val"}:
            candidate.classification = candidate.classification or "分類A"
            candidate.quality = candidate.quality or "良"
            if not candidate.mask_revisions:
                candidate.mask_revisions = ["rev_001"]
                candidate.selected_mask_revision = "rev_001"
    item = next(item for item in shell.ctx.backend.get_working_items() if item.usage == "train")
    shell.ctx.backend.update_item("all", item.item_id, usage="val")
    shell.ctx.backend.finalize_working_dataset("検証用版を追加")
    window = shell.manager.window(ModeId.COMPARISON)
    assert window.tabs.currentIndex() == 0
    QTest.mouseClick(shell.home.pipeline._stages[2], Qt.MouseButton.LeftButton)
    qapp.processEvents()
    assert window.tabs.currentIndex() == 0
    assert page.validation.findText("val_v004") >= 0


def test_rereview_01_window_activation_refreshes_the_active_page(shell, qapp):
    shell.navigate(PageId.CANDIDATES)
    page = shell.page(PageId.CANDIDATES)
    window = shell.manager.window(ModeId.COMPARISON)
    for item in shell.ctx.backend.get_working_items():
        if item.usage in {"train", "val"}:
            item.classification = item.classification or "分類A"
            item.quality = item.quality or "良"
            if not item.mask_revisions:
                item.mask_revisions = ["rev_001"]
                item.selected_mask_revision = "rev_001"
    train_item = next(
        item for item in shell.ctx.backend.get_working_items() if item.usage == "train"
    )
    shell.ctx.backend.update_item("all", train_item.item_id, usage="val")
    shell.ctx.backend.finalize_working_dataset("ウィンドウ再表示の確認")
    assert page.validation.findText("val_v004") < 0
    QTest.qWait(550)
    qapp.sendEvent(window, QEvent(QEvent.Type.WindowDeactivate))
    qapp.sendEvent(window, QEvent(QEvent.Type.WindowActivate))
    qapp.processEvents()
    assert page.validation.findText("val_v004") >= 0


def test_review_07_empty_excel_export_writes_header_without_slot_exception(
    shell, qapp, tmp_path, monkeypatch
):
    pytest.importorskip("openpyxl")
    page = shell.page(PageId.DATA_PREPARATION)
    page.search.setText("no-such-item")
    output = str(tmp_path / "empty.xlsx")
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *_args: (output, "Excel"))
    _click_menu_action(qapp, page.excel_button, page.excel_button.menu().actions()[0])
    from openpyxl import load_workbook

    workbook = load_workbook(output, read_only=True)
    try:
        assert workbook.active.max_row == 1
    finally:
        workbook.close()


def test_review_07_data_menu_actions_live_as_long_as_the_page(shell):
    page = shell.page(PageId.DATA_PREPARATION)
    assert page.mask_revision_action.parent() is page
    assert page.archive_action.parent() is page
    page.refresh()


def test_review_08_error_chip_uses_all_errors_after_filtering(shell, qapp):
    page = shell.page(PageId.DATA_PREPARATION)
    expected = len(shell.ctx.backend.validate_items().errors)
    assert int(page.chips["errors"].count_label.text()) == expected
    QTest.mouseClick(page.chips["errors"], Qt.MouseButton.LeftButton)
    assert page.model.rowCount() == expected
    assert page.finalize_button.isEnabled() == (expected == 0)


def test_review_09_keymap_cancel_discards_staged_changes(shell, qapp):
    shortcuts = shell.ctx.shortcuts
    original = shortcuts["usage_train"]
    dialog = KeymapDialog(None, shortcuts)
    row = dialog.actions.index("usage_train")
    dialog.table.setCurrentCell(row, 2)
    dialog.table.setFocus()
    QTest.keyClick(dialog.table, Qt.Key.Key_F2)
    editor = dialog.table.focusWidget()
    assert isinstance(editor, ShortcutKeyEdit)
    editor.setFocus()
    QTest.keyClick(editor, Qt.Key.Key_T)
    QTest.keyClick(editor, Qt.Key.Key_Return)
    qapp.processEvents()
    assert dialog.shortcuts["usage_train"] == "T"
    dialog.reject()
    assert shortcuts["usage_train"] == original


def test_review_10_keymap_editor_captures_modifier_sequence(qapp, shell):
    dialog = KeymapDialog(None, shell.ctx.shortcuts)
    row = dialog.actions.index("display_mode")
    dialog.table.setCurrentCell(row, 2)
    dialog.table.setFocus()
    QTest.keyClick(dialog.table, Qt.Key.Key_F2)
    editor = dialog.table.focusWidget()
    assert isinstance(editor, ShortcutKeyEdit)
    editor.setFocus()
    QTest.keyClick(
        editor,
        Qt.Key.Key_J,
        Qt.KeyboardModifier.ControlModifier | Qt.KeyboardModifier.AltModifier,
    )
    QTest.keyClick(editor, Qt.Key.Key_Return)
    qapp.processEvents()
    assert dialog.table.item(row, 2).text() == "Ctrl+Alt+J"
    assert dialog.shortcuts["display_mode"] == "Ctrl+Alt+J"
    dialog.reject()


def test_review_11_saving_keymap_by_mouse_updates_open_pages_and_triage(qapp, shell):
    data_page = shell.page(PageId.DATA_PREPARATION)
    compare_page = shell.page(PageId.MASK_COMPARISON)
    inference_page = shell.page(PageId.INFERENCE)
    dialog_results = []

    def edit_and_save(dialog):
        dialog.show()
        qapp.processEvents()
        for action_name, key, modifiers in (
            ("usage_train", Qt.Key.Key_T, Qt.KeyboardModifier.AltModifier),
            ("display_mode", Qt.Key.Key_U, Qt.KeyboardModifier.NoModifier),
        ):
            row = dialog.actions.index(action_name)
            cell = dialog.table.item(row, 2)
            dialog.table.scrollToItem(cell)
            qapp.processEvents()
            rect = dialog.table.visualItemRect(cell)
            QTest.mouseClick(dialog.table.viewport(), Qt.MouseButton.LeftButton, pos=rect.center())
            QTest.keyClick(dialog.table, Qt.Key.Key_F2)
            qapp.processEvents()
            editor = dialog.table.focusWidget()
            assert isinstance(editor, ShortcutKeyEdit)
            QTest.keyClick(editor, key, modifiers)
            qapp.processEvents()
        QTest.mouseClick(
            dialog.buttons.button(QDialogButtonBox.StandardButton.Save),
            Qt.MouseButton.LeftButton,
        )
        qapp.processEvents()
        dialog_results.append(dialog.result())
        return dialog.result()

    action = next(
        action
        for action in data_page.other_button.menu().actions()
        if action.text() == "キー割り当て一覧…"
    )
    with patch.object(KeymapDialog, "exec", edit_and_save):
        _click_menu_action(qapp, data_page.other_button, action)
        shell.ctx.keymap_window.change_button.click()
    assert dialog_results == [QDialog.DialogCode.Accepted]
    assert shell.ctx.shortcuts["usage_train"] == "Alt+T"
    assert shell.ctx.shortcuts["display_mode"] == "U"
    assert (
        shell.ctx.keymap_window.table.item(
            shell.ctx.keymap_window.actions.index("display_mode"), 2
        ).text()
        == "U"
    )
    assert compare_page.shortcuts["display_mode"] == "U"
    assert inference_page.shortcuts["usage_train"] == "Alt+T"
    item = next(item for item in data_page.items if item.usage == "unassigned")
    dialog = ContinuousTriageDialog(
        data_page,
        [item],
        lambda item_id, **changes: shell.ctx.backend.update_item("all", item_id, **changes),
        shell.ctx.backend,
        shell.ctx.shortcuts,
    )
    assert "Alt+T" in dialog.usage_buttons["train"][0].text()
    dialog.show()
    dialog.setFocus()
    QTest.keyClick(dialog, Qt.Key.Key_T, Qt.KeyboardModifier.AltModifier)
    qapp.processEvents()
    assert item.usage == "train"
    dialog.close()


def test_rereview_02_shifted_punctuation_shortcuts_and_editor_normalize(qapp, shell):
    page = shell.page(PageId.DATA_PREPARATION)
    selected = page.model.visible_items()[0]
    _click_table_row(qapp, page, selected.item_id)
    page.table.setFocus()
    original_usage = selected.usage
    QTest.keyClick(page.table, Qt.Key.Key_Q, Qt.KeyboardModifier.ShiftModifier)
    assert selected.usage == original_usage
    before = page.image_view.zoom
    QTest.keyClick(page.table, Qt.Key.Key_Equal, Qt.KeyboardModifier.ShiftModifier)
    qapp.processEvents()
    assert page.image_view.zoom > before
    with patch.object(page, "show_key_help") as show_help:
        QTest.keyClick(page.table, Qt.Key.Key_Slash, Qt.KeyboardModifier.ShiftModifier)
        qapp.processEvents()
    show_help.assert_called_once()

    dialog = KeymapDialog(None, shell.ctx.shortcuts)
    row = dialog.actions.index("zoom_out")
    cell = dialog.table.item(row, 2)
    dialog.show()
    dialog.table.scrollToItem(cell)
    qapp.processEvents()
    rect = dialog.table.visualItemRect(cell)
    QTest.mouseClick(dialog.table.viewport(), Qt.MouseButton.LeftButton, pos=rect.center())
    QTest.keyClick(dialog.table, Qt.Key.Key_F2)
    qapp.processEvents()
    editor = dialog.table.focusWidget()
    assert isinstance(editor, ShortcutKeyEdit)
    QTest.keyClick(editor, Qt.Key.Key_Minus, Qt.KeyboardModifier.ShiftModifier)
    qapp.processEvents()
    assert dialog.shortcuts["zoom_out"] == "-"
    QTest.mouseClick(
        dialog.buttons.button(QDialogButtonBox.StandardButton.Save), Qt.MouseButton.LeftButton
    )
    qapp.processEvents()
    assert shell.ctx.shortcuts["zoom_out"] == "-"


def test_rereview_02_keymap_capture_normalizes_shifted_plus(qapp, shell):
    dialog = KeymapDialog(None, shell.ctx.shortcuts)
    row = dialog.actions.index("zoom_in")
    cell = dialog.table.item(row, 2)
    dialog.show()
    dialog.table.scrollToItem(cell)
    qapp.processEvents()
    rect = dialog.table.visualItemRect(cell)
    QTest.mouseClick(dialog.table.viewport(), Qt.MouseButton.LeftButton, pos=rect.center())
    QTest.keyClick(dialog.table, Qt.Key.Key_F2)
    qapp.processEvents()
    editor = dialog.table.focusWidget()
    assert isinstance(editor, ShortcutKeyEdit)
    QTest.keyClick(editor, Qt.Key.Key_Equal, Qt.KeyboardModifier.ShiftModifier)
    qapp.processEvents()
    assert dialog.shortcuts["zoom_in"] == "+"
    QTest.mouseClick(
        dialog.buttons.button(QDialogButtonBox.StandardButton.Save), Qt.MouseButton.LeftButton
    )
    qapp.processEvents()
    assert shell.ctx.shortcuts["zoom_in"] == "+"


def test_review_12_default_shortcuts_cover_clear_help_navigation_and_zoom(shell, qapp, monkeypatch):
    shell.navigate(PageId.DATA_PREPARATION)
    page = shell.page(PageId.DATA_PREPARATION)
    item = next(item for item in page.items if item.usage == "unassigned")
    _click_table_row(qapp, page, item.item_id)
    page.table.setFocus()
    QTest.keyClick(page.table, Qt.Key.Key_0)
    assert item.classification is None
    with patch(
        "foam_cell_analysis.gui.modes.data_preparation.page.QInputDialog.getItem",
        return_value=("分類B", True),
    ):
        QTest.keyClick(page.table, Qt.Key.Key_K)
    assert item.classification == "分類B"
    QTest.keyClick(page.table, Qt.Key.Key_F1)
    assert shell.ctx.keymap_window is not None
    assert shell.ctx.keymap_window.isVisible()
    before = page.image_view.zoom
    QTest.keyClick(page.image_view, Qt.Key.Key_Plus)
    assert page.image_view.zoom > before
    QTest.keyClick(page.image_view, Qt.Key.Key_F)
    assert page.image_view.zoom <= before
    QTest.keyClick(page.table, Qt.Key.Key_Escape)
    assert not page._selected_ids
    QTest.keyClick(page.table, Qt.Key.Key_Escape)
    assert page.model.usages is None
    QTest.keyClick(page.table, Qt.Key.Key_A, Qt.KeyboardModifier.ControlModifier)
    assert len(page.table.selectionModel().selectedRows()) == page.model.rowCount()
    all_chip_count = page.model.rowCount()
    QTest.keyClick(page.table, Qt.Key.Key_6, Qt.KeyboardModifier.AltModifier)
    qapp.processEvents()
    assert page.model.errors_only
    assert page.model.rowCount() == int(page.chips["errors"].count_label.text())
    assert all_chip_count >= page.model.rowCount()


def test_review_12_shift_space_selects_previous_unassigned(qapp, shell):
    page = shell.page(PageId.DATA_PREPARATION)
    unassigned = [item for item in page.model.visible_items() if item.usage == "unassigned"]
    assert len(unassigned) >= 2
    _click_table_row(qapp, page, unassigned[0].item_id)
    page.table.setFocus()
    QTest.keyClick(page.table, Qt.Key.Key_Space, Qt.KeyboardModifier.ShiftModifier)
    assert page._selected_ids == [unassigned[-1].item_id]
    QTest.keyClick(page.table, Qt.Key.Key_Space)
    assert page._selected_ids == [unassigned[0].item_id]


def test_review_13_candidate_menu_actions_disable_without_selection(qapp, qtbot):
    page = CandidatesPage(AppContext(MockBackend(), Navigator(), JobManager()))
    qtbot.addWidget(page)
    page.show()
    assert not page.menu_actions["export"].isEnabled()
    assert not page.menu_actions["reject"].isEnabled()
    page.menu_actions["export"].trigger()
    assert not page.menu_actions["export"].isEnabled()


def test_review_14_imported_usage_is_visible_and_source_filter_says_all(shell, qapp):
    page = shell.page(PageId.DATA_PREPARATION)

    def finish_import(dialog):
        dialog.image_edit.setText("C:/review/images")
        dialog.mask_edit.setText("C:/review/masks")
        dialog._accept()
        return dialog.result()

    def finish_settings(dialog):
        for _folder, usage, classification, quality in dialog.rows:
            usage.setCurrentIndex(usage.findData("train"))
            classification.setCurrentIndex(classification.findData("分類A"))
            quality.setCurrentIndex(quality.findData("良"))
        dialog._accept()
        return dialog.result()

    with (
        patch.object(ImportDialog, "exec", finish_import),
        patch.object(ImportSettingsDialog, "exec", finish_settings),
    ):
        QTest.mouseClick(page.import_button, Qt.MouseButton.LeftButton)
        qapp.processEvents()
    assert page.model.rowCount() > 0
    assert page.model.usages == {"train"}
    assert page.source_combo.currentData() is None


def test_review_15_reopening_current_mode_from_home_keeps_active_tab(shell, qapp):
    shell.navigate(PageId.EXPERIMENTS)
    window = shell.manager.window(ModeId.TRAINING)
    QTest.mouseClick(shell.home.pipeline._stages[1], Qt.MouseButton.LeftButton)
    qapp.processEvents()
    assert window.tabs.currentIndex() == 2
    assert shell.manager.current_page_id(ModeId.TRAINING) == PageId.EXPERIMENTS


def test_review_16_home_refreshes_progress_from_started_training(shell, qapp, monkeypatch):
    from foam_cell_analysis.gui.modes.training import page as training_module

    real_fake_job = training_module.FakeJob

    def quick_job(*args, **kwargs):
        kwargs["interval_ms"] = 20
        return real_fake_job(*args, **kwargs)

    monkeypatch.setattr(training_module, "FakeJob", quick_job)
    shell.navigate(PageId.TRAINING)
    page = shell.page(PageId.TRAINING)
    with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
        QTest.mouseClick(page.start_button, Qt.MouseButton.LeftButton)
    qapp.processEvents()
    assert shell.ctx.jobs.running_count == 1
    QTest.qWait(60)
    qapp.processEvents()
    assert shell.home._summary.running_experiment
    assert shell.home._summary.running_epoch >= 1
    assert shell.home.pipeline._values[1].text() == str(shell.home._summary.running_epoch)
    shell.ctx.jobs.jobs()[0].cancel()


def test_review_17_triage_returns_current_item_and_stops_at_exhaustion(shell, qapp, qtbot):
    page = shell.page(PageId.DATA_PREPARATION)
    targets = [item for item in page.items if item.usage == "unassigned"][:2]
    dialog = ContinuousTriageDialog(
        page,
        targets,
        lambda item_id, **changes: shell.ctx.backend.update_item("all", item_id, **changes),
        shell.ctx.backend,
        shell.ctx.shortcuts,
    )
    qtbot.addWidget(dialog)
    dialog.show()
    rect = dialog.filmstrip.visualItemRect(dialog.filmstrip.item(1))
    QTest.mouseClick(dialog.filmstrip.viewport(), Qt.MouseButton.LeftButton, pos=rect.center())
    assert dialog.current_item_id == targets[1].item_id
    dialog.setFocus()
    QTest.keyClick(dialog, Qt.Key.Key_Q)
    QTest.keyClick(dialog, Qt.Key.Key_Q)
    qapp.processEvents()
    assert dialog.filmstrip.count() == 0
    assert "ありません" in dialog.counter.text()
    dialog.close()


def test_review_18_excel_import_releases_source_file(shell, qapp, tmp_path, monkeypatch):
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    page = shell.page(PageId.DATA_PREPARATION)
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["id", "folder", "filename", "usage", "class", "quality", "mask"])
    item = page.items[0]
    sheet.append([item.item_id, "folder", item.source_filename, "学習", "分類A", "良", "rev_001"])
    workbook.save(path)
    workbook.close()
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *_args: (str(path), "Excel"))
    _click_menu_action(qapp, page.excel_button, page.excel_button.menu().actions()[1])
    path.rename(tmp_path / "renamed.xlsx")
    assert (tmp_path / "renamed.xlsx").exists()


def test_review_19_data_preparation_other_menu_has_required_actions(shell):
    page = shell.page(PageId.DATA_PREPARATION)
    labels = [action.text() for action in page.other_button.menu().actions()]
    assert "新しいマスク版を取り込む…" in labels
    assert "アーカイブを作成…" in labels
    assert "キー割り当て一覧…" in labels


def test_review_19_archive_menu_action_records_archive_for_latest_versions(shell, monkeypatch):
    shell.page(PageId.DATA_PREPARATION)
    monkeypatch.setattr(QFileDialog, "getExistingDirectory", lambda *_args: "C:/archive")
    window = shell.manager.window(ModeId.DATA_PREPARATION)
    file_menu = next(
        action.menu()
        for action in window.menuBar().actions()
        if action.text().startswith("ファイル")
    )
    action = next(action for action in file_menu.actions() if action.text() == "アーカイブを作成…")
    assert action.isEnabled()
    action.trigger()
    latest = {}
    for version in shell.ctx.backend.list_dataset_versions():
        latest[version.purpose] = version
    assert all(version.archive_status == "COMPLETED" for version in latest.values())


def test_review_19_mask_import_menu_action_adds_revision_to_selected_item(shell):
    page = shell.page(PageId.DATA_PREPARATION)
    item = next(item for item in page.items if item.item_id in page._selected_ids)
    previous_revisions = len(item.mask_revisions)
    window = shell.manager.window(ModeId.DATA_PREPARATION)
    file_menu = next(
        action.menu()
        for action in window.menuBar().actions()
        if action.text().startswith("ファイル")
    )
    action = next(
        action for action in file_menu.actions() if action.text() == "新しいマスク版を取り込む…"
    )
    assert action.isEnabled()
    action.trigger()
    assert len(item.mask_revisions) == previous_revisions + 1


def test_review_20_thumbnail_tab_counts_show_filtered_and_total(qapp):
    backend = MockBackend()
    train_items = [item for item in backend.get_working_items() if item.usage == "train"]
    train_items[0].change = None
    dialog = DatasetFinalizeDialog(None, backend)
    dialog.show()
    tab_bar = dialog.finalize_tabs.tabBar()
    QTest.mouseClick(tab_bar, Qt.MouseButton.LeftButton, pos=tab_bar.tabRect(1).center())
    qapp.processEvents()
    _choose_combo_option(qapp, dialog.thumbnail_filter, "すべて")
    _choose_combo_option(qapp, dialog.thumbnail_filter, "追加・変更のみ")
    for purpose, label in (("train", "学習用"), ("val", "検証用")):
        model = dialog.thumbnail_models[purpose]
        index = dialog.thumbnail_tabs.indexOf(dialog.thumbnail_stacks[purpose])
        assert dialog.thumbnail_tabs.tabText(index) == (
            f"{label} ({model.rowCount()} / {len(dialog.thumbnail_source_items[purpose])} 件)"
        )


def test_review_21_empty_thumbnail_filter_shows_centered_guidance(qapp):
    backend = MockBackend()
    for item in backend.get_working_items():
        item.change = None
    dialog = DatasetFinalizeDialog(None, backend)
    dialog.show()
    tab_bar = dialog.finalize_tabs.tabBar()
    QTest.mouseClick(tab_bar, Qt.MouseButton.LeftButton, pos=tab_bar.tabRect(1).center())
    qapp.processEvents()
    _choose_combo_option(qapp, dialog.thumbnail_filter, "すべて")
    assert any(model.rowCount() for model in dialog.thumbnail_models.values())
    _choose_combo_option(qapp, dialog.thumbnail_filter, "追加・変更のみ")
    for label in dialog.thumbnail_empty_labels.values():
        assert label.text() == "追加・変更された画像はありません"
        assert label.alignment() & Qt.AlignmentFlag.AlignHCenter
        assert (
            dialog.thumbnail_stacks[
                next(
                    purpose
                    for purpose, target in dialog.thumbnail_empty_labels.items()
                    if target is label
                )
            ].currentWidget()
            is label
        )


def test_rereview_03_data_preview_zoom_survives_activation_after_table_click(shell, qapp):
    page = shell.page(PageId.DATA_PREPARATION)
    item = page.model.visible_items()[0]
    _click_table_row(qapp, page, item.item_id)
    page.table.setFocus()
    QTest.keyClick(page.table, Qt.Key.Key_Equal, Qt.KeyboardModifier.ShiftModifier)
    qapp.processEvents()
    zoom_before = page.image_view.zoom
    window = shell.manager.window(ModeId.DATA_PREPARATION)
    QTest.qWait(550)
    qapp.sendEvent(window, QEvent(QEvent.Type.WindowDeactivate))
    qapp.sendEvent(window, QEvent(QEvent.Type.WindowActivate))
    qapp.processEvents()
    assert page.image_view.zoom == pytest.approx(zoom_before)
    assert page._selected_ids == [item.item_id]


def test_rereview_04_candidate_current_row_scroll_and_space_survive_activation(shell, qapp):
    import copy

    backend = shell.ctx.backend
    base = next(iter(backend.candidates.values()))
    for number in range(4, 40):
        candidate = copy.deepcopy(base)
        candidate.candidate_id = f"ACT-{number:03}"
        backend.candidates[candidate.candidate_id] = candidate
    shell.navigate(PageId.CANDIDATES)
    page = shell.page(PageId.CANDIDATES)
    window = shell.manager.window(ModeId.COMPARISON)
    page.table.verticalScrollBar().setValue(page.table.verticalScrollBar().maximum())
    qapp.processEvents()
    row = page.table.rowCount() - 3
    candidate_id = page.table.item(row, 1).text()
    rect = page.table.visualItemRect(page.table.item(row, 0))
    QTest.mouseClick(
        page.table.viewport(),
        Qt.MouseButton.LeftButton,
        pos=rect.topLeft() + QPoint(12, rect.height() // 2),
    )
    qapp.processEvents()
    scroll_before = page.table.verticalScrollBar().value()
    assert page.table.item(row, 0).checkState() == Qt.CheckState.Checked
    QTest.qWait(550)
    qapp.sendEvent(window, QEvent(QEvent.Type.WindowDeactivate))
    qapp.sendEvent(window, QEvent(QEvent.Type.WindowActivate))
    qapp.processEvents()
    current_row = page.table.currentRow()
    assert current_row >= 0
    assert page.table.item(current_row, 1).text() == candidate_id
    assert page.table.verticalScrollBar().value() == scroll_before
    page.table.setFocus()
    QTest.keyClick(page.table, Qt.Key.Key_Space)
    qapp.processEvents()
    assert page.table.item(current_row, 0).checkState() == Qt.CheckState.Unchecked


def test_data_preparation_selection_change_fits_new_preview(shell, qapp):
    page = shell.page(PageId.DATA_PREPARATION)
    items = page.model.visible_items()
    _click_table_row(qapp, page, items[0].item_id)
    fitted_zoom = page.image_view.zoom
    page.image_view.zoom_by(1.5)
    assert page.image_view.zoom > fitted_zoom

    _click_table_row(qapp, page, items[1].item_id)
    QTest.qWait(10)
    qapp.processEvents()

    assert page.image_view.zoom == pytest.approx(fitted_zoom)


def test_multi_selection_survives_activation_and_reaches_auto_triage(shell, qapp):
    """複数選択 → ウィンドウの前面化 → 自動振り分けで、選択した全件が対象になる。"""
    from PySide6.QtWidgets import QApplication, QPushButton

    from foam_cell_analysis.gui.modes.data_preparation.dialogs import AutoTriageDialog

    shell.navigate(PageId.DATA_PREPARATION)
    page = shell.page(PageId.DATA_PREPARATION)
    window = shell.manager.window(ModeId.DATA_PREPARATION)
    window.resize(1400, 900)
    window.show()
    qapp.processEvents()
    viewport = page.table.viewport()
    first = page.table.visualRect(page.model.index(3, 1)).center()
    last = page.table.visualRect(page.model.index(10, 1)).center()
    QTest.mouseClick(viewport, Qt.MouseButton.LeftButton, pos=first)
    QTest.mouseClick(
        viewport, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.ShiftModifier, pos=last
    )
    selected = list(page._selected_ids)
    assert len(selected) == 8
    QTest.qWait(550)
    qapp.sendEvent(window, QEvent(QEvent.Type.WindowDeactivate))
    qapp.sendEvent(window, QEvent(QEvent.Type.WindowActivate))
    qapp.processEvents()
    assert page._selected_ids == selected
    assert len(page.table.selectionModel().selectedRows()) == 8
    seen = {}

    def run():
        dialog = QApplication.activeModalWidget()
        assert isinstance(dialog, AutoTriageDialog)
        seen["label"] = dialog.target_selected.text()
        QTest.mouseClick(dialog.target_selected, Qt.MouseButton.LeftButton)
        dialog.ratio.setValue(100)
        button = next(b for b in dialog.findChildren(QPushButton) if "振り分ける" in b.text())
        QTest.mouseClick(button, Qt.MouseButton.LeftButton)

    QTimer.singleShot(50, run)
    QTest.mouseClick(page.auto_button, Qt.MouseButton.LeftButton)
    qapp.processEvents()
    assert seen["label"] == "選択中（8件）"
    usages = {item.item_id: item.usage for item in shell.ctx.backend.get_working_items()}
    assert all(usages[item_id] == "val" for item_id in selected)


def _activate_mode_window(qapp, window):
    """WindowManager 管理下のモード窓へ前面化イベントを送る。"""
    QTest.qWait(550)
    qapp.sendEvent(window, QEvent(QEvent.Type.WindowDeactivate))
    qapp.sendEvent(window, QEvent(QEvent.Type.WindowActivate))
    qapp.processEvents()


def test_review_20_dataset_history_row_click_opens_thumbnails_and_survives_activation(shell, qapp):
    shell.navigate(PageId.DATASET_HISTORY)
    page = shell.page(PageId.DATASET_HISTORY)
    window = shell.manager.window(ModeId.DATA_PREPARATION)
    window.show()
    qapp.processEvents()
    assert page.model.rowCount() > 0
    index = page.model.index(0, 0)
    QTest.mouseClick(
        page.table.viewport(),
        Qt.MouseButton.LeftButton,
        pos=page.table.visualRect(index).center(),
    )
    qapp.processEvents()
    selected_version = page.model.versions[0].version
    assert [row.row() for row in page.table.selectionModel().selectedRows()] == [0]
    QTest.mouseClick(page.thumbnail_button, Qt.MouseButton.LeftButton)
    qapp.processEvents()
    assert selected_version in page.thumbnail_windows
    _activate_mode_window(qapp, window)
    assert [row.row() for row in page.table.selectionModel().selectedRows()] == [0]
    assert page.table.currentIndex().row() == 0


def test_review_21_triage_close_restores_multi_selection_and_current_image(shell, qapp):
    shell.navigate(PageId.DATA_PREPARATION)
    page = shell.page(PageId.DATA_PREPARATION)
    window = shell.manager.window(ModeId.DATA_PREPARATION)
    window.show()
    qapp.processEvents()
    page.chips["all"].click()
    qapp.processEvents()
    visible = page.model.visible_items()
    chosen = [visible[row].item_id for row in (3, 5, 8)]
    for position, item_id in enumerate(chosen):
        _click_table_row(
            qapp,
            page,
            item_id,
            Qt.KeyboardModifier.NoModifier
            if position == 0
            else Qt.KeyboardModifier.ControlModifier,
        )
    _activate_mode_window(qapp, window)
    assert set(page._selected_ids) == set(chosen)
    dialog_state = {}

    def close_when_open():
        dialog = QApplication.activeModalWidget()
        if not isinstance(dialog, ContinuousTriageDialog):
            QTimer.singleShot(25, close_when_open)
            return
        dialog_state["current_id"] = dialog.current_item_id
        QTest.keyClick(dialog, Qt.Key.Key_Return)

    QTimer.singleShot(0, close_when_open)
    page.table.setFocus()
    QTest.keyClick(page.table, Qt.Key.Key_Return)
    qapp.processEvents()
    visible_after = page.model.visible_items()
    selected_after = [index.row() for index in page.table.selectionModel().selectedRows()]
    assert set(page._selected_ids) == set(chosen)
    assert "current_id" in dialog_state
    assert len(selected_after) == 3
    assert page.table.currentIndex().row() == next(
        row for row, item in enumerate(visible_after) if item.item_id == dialog_state["current_id"]
    )


def test_review_22_inference_multi_selection_survives_activation(shell, qapp):
    shell.navigate(PageId.INFERENCE)
    page = shell.page(PageId.INFERENCE)
    window = shell.manager.window(ModeId.INFERENCE)
    window.show()
    page.inputs = [InferenceInput(f"{index}.png", "分類A", "model_007") for index in range(5)]
    page._refresh_table()
    qapp.processEvents()
    for row, modifiers in (
        (0, Qt.KeyboardModifier.NoModifier),
        (3, Qt.KeyboardModifier.ControlModifier),
    ):
        index = page.table.model().index(row, 0)
        QTest.mouseClick(
            page.table.viewport(),
            Qt.MouseButton.LeftButton,
            modifiers,
            page.table.visualRect(index).center(),
        )
    _activate_mode_window(qapp, window)
    selected = {index.row() for index in page.table.selectionModel().selectedRows()}
    assert selected == {0, 3}
    assert page.table.currentRow() == 3


def test_review_23_released_model_multi_selection_survives_activation(shell, qapp):
    shell.navigate(PageId.RELEASED_MODELS)
    page = shell.page(PageId.RELEASED_MODELS)
    window = shell.manager.window(ModeId.COMPARISON)
    window.show()
    qapp.processEvents()
    assert page.model_table.rowCount() >= 2
    for row, modifiers in (
        (0, Qt.KeyboardModifier.NoModifier),
        (1, Qt.KeyboardModifier.ControlModifier),
    ):
        index = page.model_table.model().index(row, 0)
        QTest.mouseClick(
            page.model_table.viewport(),
            Qt.MouseButton.LeftButton,
            modifiers,
            page.model_table.visualRect(index).center(),
        )
    _activate_mode_window(qapp, window)
    selected = {index.row() for index in page.model_table.selectionModel().selectedRows()}
    assert selected == {0, 1}
    assert page.model_table.currentRow() == 1
