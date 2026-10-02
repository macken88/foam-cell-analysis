"""データ準備 v2 の一括作業フローを検証する。"""

from __future__ import annotations

from unittest.mock import patch

import pytest
from PySide6.QtCore import QEvent, QPoint, Qt, QTimer
from PySide6.QtTest import QSignalSpy, QTest
from PySide6.QtWidgets import QApplication, QFileDialog, QSizePolicy

from foam_cell_analysis.gui.context import AppContext, StatusBus
from foam_cell_analysis.gui.jobs import JobManager
from foam_cell_analysis.gui.modes.data_preparation.dialogs import (
    AutoTriageDialog,
    ContinuousTriageDialog,
    DatasetFinalizeDialog,
    ImportDialog,
    ImportSettingsDialog,
)
from foam_cell_analysis.gui.modes.data_preparation.page import (
    DataPreparationPage,
    DatasetHistoryPage,
)
from foam_cell_analysis.gui.modes.training.page import TrainingPage
from foam_cell_analysis.gui.navigation import ModeId, Navigator, PageId
from foam_cell_analysis.gui.theme import numeric_font
from foam_cell_analysis.gui.widgets.image_convert import DisplayMode
from foam_cell_analysis.services.mock.backend import MockBackend
from foam_cell_analysis.services.models import DataItem


def _page(qapp, backend=None):
    backend = backend or MockBackend()
    return DataPreparationPage(AppContext(backend, Navigator(), JobManager()))


def _item(backend, item_id: str, folder: str, usage: str = "unassigned", classification="分類A"):
    return DataItem(
        item_id,
        f"{item_id}.tif",
        f"{folder}/{item_id}.tif",
        ["A"],
        classification,
        "良",
        ["rev_001"],
        "rev_001",
        usage=usage,
        sha256=backend._digest(item_id),
        seed=int(backend._digest(item_id)[:6], 16),
    )


def test_import_uniform_values_are_applied_per_item_without_linking_edits(qapp):
    backend = MockBackend()
    candidates = backend.scan_import_source({"A": "C:/batch-a"}, "C:/masks")[:3]
    groups = {candidates[0].source_relpath.rsplit("/", 1)[0]: candidates}
    dialog = ImportSettingsDialog(None, groups)
    dialog.rows[0][1].setCurrentIndex(dialog.rows[0][1].findData("train"))
    dialog.rows[0][2].setCurrentIndex(dialog.rows[0][2].findData("分類B"))
    dialog.rows[0][3].setCurrentIndex(dialog.rows[0][3].findData("良"))
    dialog._accept()
    values = dialog.values[next(iter(groups))]
    items = backend.import_items("all", candidates, values["classification"])
    backend.bulk_update_items("all", [item.item_id for item in items], **values)
    first, second = items[:2]
    backend.update_item("all", first.item_id, classification="分類C")
    assert first.classification == "分類C"
    assert second.classification == "分類B"
    assert all(item.usage == "train" and item.quality == "良" for item in items)


def test_table_edits_multiple_items_and_undo_restores_them(qapp):
    backend = MockBackend()
    page = _page(qapp, backend)
    selected = [item.item_id for item in backend.get_working_items()[:2]]
    page._selected_ids = selected
    page._change(selected, usage="val")
    assert all(
        next(item for item in backend.get_working_items() if item.item_id == key).usage == "val"
        for key in selected
    )
    page.undo_stack.undo()
    assert all(
        next(item for item in backend.get_working_items() if item.item_id == key).usage == "train"
        for key in selected
    )
    row = page.model.visible_items().index(backend.get_working_items()[0])
    assert page.model.setData(page.model.index(row, 4), "可", Qt.ItemDataRole.EditRole)
    assert backend.get_working_items()[0].quality == "可"
    page.undo_stack.undo()
    assert backend.get_working_items()[0].quality == "良"
    page.close()


def _visible_ids(page):
    return [item.item_id for item in page.model.visible_items()]


def _click_combo_row(qtbot, combo, row):
    combo.showPopup()
    qtbot.wait(20)
    index = combo.model().index(row, 0)
    rect = combo.view().visualRect(index)
    QTest.mouseClick(
        combo.view().viewport(),
        Qt.MouseButton.LeftButton,
        pos=rect.topLeft() + QPoint(10, rect.height() // 2),
    )
    qtbot.wait(20)


def test_filter_controls_combine_and_match_visible_rows(shell, qtbot):
    from foam_cell_analysis.gui.navigation import PageId

    shell.navigate(PageId.DATA_PREPARATION)
    page = shell.page(PageId.DATA_PREPARATION)
    window = shell.manager.window(
        __import__("foam_cell_analysis.gui.navigation", fromlist=["ModeId"]).ModeId.DATA_PREPARATION
    )
    window.show()
    window.activateWindow()
    window.raise_()
    qtbot.waitUntil(window.isActiveWindow)
    qtbot.waitUntil(lambda: page.isVisible())

    first_usage = page.items[0].usage
    if first_usage != "train":
        target_usage = next((item for item in page.items if item.usage == "train"), None)
    else:
        target_usage = page.items[0]
    assert target_usage is not None
    target_usage.change = "updated"
    target_usage.classification = "分類A"
    target_usage.source_relpath = "test-source/needle.tif"
    target_usage.source_filename = "needle.tif"
    second_target = next(
        item
        for item in page.items
        if item.item_id != target_usage.item_id and item.usage == "train"
    )
    second_target.change = "updated"
    second_target.classification = "分類B"
    second_target.source_relpath = "test-source/other.tif"
    second_target.source_filename = "other.tif"
    shell.ctx.backend.validate_items = lambda *_args, **_kwargs: type(
        "Report",
        (),
        {
            "errors": [
                type("Issue", (), {"item_id": item.item_id, "message": "test error"})()
                for item in (target_usage, second_target)
            ]
        },
    )()
    page.refresh()

    QTest.mouseClick(page.chips["train"], Qt.MouseButton.LeftButton)
    QTest.mouseClick(page.error_filter, Qt.MouseButton.LeftButton)
    QTest.mouseClick(page.changed_filter, Qt.MouseButton.LeftButton)
    expected = {
        item.item_id
        for item in page.items
        if item.usage == "train" and item.item_id in page.model.errors and item.change is not None
    }
    assert set(_visible_ids(page)) == expected
    _click_combo_row(qtbot, page.class_filter, 1)
    _click_combo_row(qtbot, page.class_filter, 2)
    assert page.model.classifications == {"分類A", "分類B"}
    assert set(_visible_ids(page)) == {target_usage.item_id, second_target.item_id}
    _click_combo_row(qtbot, page.class_filter, 0)
    assert page.model.classifications == set()
    before = (target_usage.usage, target_usage.classification, target_usage.quality)
    page.class_filter.setFocus()
    for key in (
        Qt.Key.Key_Q,
        Qt.Key.Key_W,
        Qt.Key.Key_E,
        Qt.Key.Key_R,
        Qt.Key.Key_1,
        Qt.Key.Key_Z,
        Qt.Key.Key_X,
        Qt.Key.Key_C,
    ):
        QTest.keyClick(page.class_filter, key)
    assert (target_usage.usage, target_usage.classification, target_usage.quality) == before
    page.source_combo.showPopup()
    qtbot.wait(20)
    source_index = page.source_combo.findData("test-source")
    QTest.mouseClick(
        page.source_combo.view().viewport(),
        Qt.MouseButton.LeftButton,
        pos=page.source_combo.view()
        .visualRect(page.source_combo.model().index(source_index, 0))
        .center(),
    )
    QTest.keyClicks(page.search, "needle")
    assert _visible_ids(page) == [target_usage.item_id]
    assert page.visible_count.text() == f"{len(page.items)} 件中 1 件を表示"

    page.search.clear()
    QTest.mouseClick(page.changed_filter, Qt.MouseButton.LeftButton)
    assert not page.model.changed_only
    page.table.setFocus()
    QTest.keyClick(page.table, Qt.Key.Key_6, Qt.KeyboardModifier.AltModifier)
    assert not page.model.errors_only
    QTest.keyClick(page.table, Qt.Key.Key_6, Qt.KeyboardModifier.AltModifier)
    assert page.model.errors_only


def test_auto_triage_is_deterministic_stratified_groupwise_and_preserves_assigned_items(qapp):
    backend = MockBackend()
    master = backend.get_working_items()
    assigned = next(item for item in master if item.usage == "train")
    existing = assigned.usage
    candidates = [
        _item(
            backend,
            f"triage_{index:03d}",
            f"lot-{index // 2}",
            classification=("分類A" if index % 2 == 0 else "分類B"),
        )
        for index in range(12)
    ]
    backend.add_imported_items("all", candidates)
    result = backend.apply_auto_triage(
        {"validation_ratio": 50, "by_folder": True, "stratify": True, "seed": 42},
        [item.item_id for item in candidates],
    )
    assert result["train"] + result["val"] == 12
    assert assigned.usage == existing
    assert all(
        len({item.usage for item in candidates if item.source_folder == folder}) == 1
        for folder in {item.source_folder for item in candidates}
    )
    after = [item.usage for item in candidates]
    for item in candidates:
        item.usage = "unassigned"
    backend.apply_auto_triage(
        {"validation_ratio": 50, "by_folder": True, "stratify": True, "seed": 42},
        [item.item_id for item in candidates],
    )
    assert [item.usage for item in candidates] == after


def test_finalization_publishes_only_changed_train_and_val_and_omits_unassigned(qapp):
    backend = MockBackend()
    master = backend.get_working_items()
    for item in master:
        if item.usage in {"train", "val"}:
            item.classification = item.classification or "分類A"
            item.quality = item.quality or "良"
            if not item.mask_revisions:
                item.mask_revisions = ["rev_001"]
                item.selected_mask_revision = "rev_001"
    candidates = backend.scan_import_source({"A": "C:/new-batch"}, "C:/masks")[:2]
    imported = backend.import_items("all", candidates, "分類A")
    extra_unassigned = _item(backend, "extra_unassigned", "lot-unassigned")
    backend.add_imported_items("all", [extra_unassigned])
    backend.bulk_update_items("all", [imported[0].item_id], usage="train", quality="良")
    backend.bulk_update_items("all", [imported[1].item_id], usage="val", quality="良")
    unassigned = next(item for item in master if item.usage == "unassigned")
    versions = backend.finalize_working_dataset("同時確定")
    assert {item.purpose for item in versions} == {"train", "val"}
    assert all(unassigned.item_id not in version.item_ids for version in versions)
    train = next(item for item in versions if item.purpose == "train")
    val = next(item for item in versions if item.purpose == "val")
    assert train.base_validation_version == val.version
    assert train.version == "train_v004" and val.version == "val_v004"


def test_excel_export_and_import_round_trip(qapp, tmp_path, monkeypatch):
    pytest = __import__("pytest")
    pytest.importorskip("openpyxl")
    backend = MockBackend()
    page = _page(qapp, backend)
    output = str(tmp_path / "working.xlsx")
    monkeypatch.setattr(
        QFileDialog, "getSaveFileName", lambda *_args, **_kwargs: (output, "Excel (*.xlsx)")
    )
    page.export_excel()
    before = [
        (item.item_id, item.usage, item.classification, item.quality)
        for item in backend.get_working_items()
    ]
    monkeypatch.setattr(
        QFileDialog, "getOpenFileName", lambda *_args, **_kwargs: (output, "Excel (*.xlsx)")
    )
    page.import_excel()
    after = [
        (item.item_id, item.usage, item.classification, item.quality)
        for item in backend.get_working_items()
    ]
    assert after == before
    page.undo_stack.undo()
    assert [
        (item.item_id, item.usage, item.classification, item.quality)
        for item in backend.get_working_items()
    ] == before
    page.close()


def test_working_table_uses_no_thumbnail_toggle_and_updates_only_edited_items(qapp):
    backend = MockBackend()
    page = _page(qapp, backend)
    assert not hasattr(page, "thumbnail")
    assert "toggle_view" not in page.shortcuts.mapping
    calls = []
    validate = backend.validate_items

    def record(item_ids=None):
        calls.append(item_ids)
        return validate(item_ids)

    backend.validate_items = record
    resets = QSignalSpy(page.model.modelReset)
    layouts = QSignalSpy(page.model.layoutChanged)
    item = next(item for item in backend.get_working_items() if item.usage == "unassigned")
    page._change([item.item_id], quality="可")
    assert calls == [None]
    assert resets.count() == 0
    assert layouts.count() == 0
    page.close()


def test_preview_shortcuts_work_from_table_but_not_search_input(qapp):
    backend = MockBackend()
    page = _page(qapp, backend)
    assert {action.text() for action in page.display_actions.values()} == {
        "正解ラベル画像",
        "正解ラベルを重ねて表示",
        "白黒の正解ラベル画像",
    }
    page.ctx.display.set_value("二値マスク")
    assert page.display_toggle.alternate_button.text() == "白黒の正解ラベル画像"
    messages = []
    page.ctx.status.message.connect(messages.append)
    page.table.setFocus()
    QTest.keyClick(page.table, Qt.Key.Key_M)
    assert not page.display_toggle.is_alternate
    QTest.keyClick(page.table, Qt.Key.Key_M)
    assert page.display_toggle.is_alternate
    page.search.setFocus()
    QTest.keyClick(page.search, Qt.Key.Key_M)
    assert page.search.text().casefold() == "m"
    assert page.display_toggle.is_alternate
    page.close()


def test_continuous_triage_shortcuts_toggle_mode_and_move_image(qapp, qtbot):
    backend = MockBackend()
    page = _page(qapp, backend)
    items = [item for item in backend.get_working_items() if item.usage == "unassigned"][:2]
    dialog = ContinuousTriageDialog(
        page, items, lambda *_args, **_kwargs: None, backend, page.shortcuts
    )
    qtbot.addWidget(dialog)
    messages = []
    page.ctx.status.message.connect(messages.append)
    dialog.show()
    dialog.setFocus()
    QTest.keyClick(dialog, Qt.Key.Key_M)
    assert dialog.display_mode == DisplayMode.IMAGE
    QTest.keyClick(dialog, Qt.Key.Key_M)
    assert dialog.display_mode == DisplayMode.OVERLAY
    QTest.keyClick(dialog, Qt.Key.Key_Right)
    assert dialog.index == 1
    dialog.close()
    page.close()


def _select_rows_by_click(page, qtbot, rows):
    page.table.clearSelection()
    for position, row in enumerate(rows):
        page.table.scrollTo(page.model.index(row, 0))
        qtbot.wait(10)
        rect = page.table.visualRect(page.model.index(row, 0))
        modifier = (
            Qt.KeyboardModifier.NoModifier if position == 0 else Qt.KeyboardModifier.ControlModifier
        )
        QTest.mouseClick(page.table.viewport(), Qt.MouseButton.LeftButton, modifier, rect.center())
    page.table.setFocus()
    qtbot.wait(10)


def test_undo_redo_shortcuts_only_apply_to_table_or_preview(shell, qapp, qtbot):
    shell.navigate(PageId.DATA_PREPARATION)
    page = shell.page(PageId.DATA_PREPARATION)
    window = shell.manager.window(ModeId.DATA_PREPARATION)
    window.show()
    window.activateWindow()
    window.raise_()
    qtbot.waitUntil(window.isActiveWindow)
    item = page.items[0]
    old_quality = item.quality
    new_quality = "可" if old_quality != "可" else "良"
    page._change([item.item_id], quality=new_quality)
    stack_index = page.undo_stack.index()

    page.source_combo.setFocus()
    qtbot.waitUntil(page.source_combo.hasFocus)
    QTest.keyClick(page.source_combo, Qt.Key.Key_Z, Qt.KeyboardModifier.ControlModifier)
    qapp.processEvents()
    assert page.undo_stack.index() == stack_index
    assert item.quality == new_quality

    page.class_filter.setFocus()
    qtbot.waitUntil(page.class_filter.hasFocus)
    QTest.keyClick(page.class_filter, Qt.Key.Key_Y, Qt.KeyboardModifier.ControlModifier)
    qapp.processEvents()
    assert page.undo_stack.index() == stack_index
    assert item.quality == new_quality

    page.table.setFocus()
    QTest.keyClick(page.table, Qt.Key.Key_Z, Qt.KeyboardModifier.ControlModifier)
    qapp.processEvents()
    assert page.undo_stack.index() == stack_index - 1
    assert item.quality == old_quality
    QTest.keyClick(page.table, Qt.Key.Key_Y, Qt.KeyboardModifier.ControlModifier)
    qapp.processEvents()
    assert page.undo_stack.index() == stack_index


def test_non_dialog_page_shortcuts_do_not_override_text_input(shell, qapp, qtbot):
    shell.navigate(PageId.DATA_PREPARATION)
    page = shell.page(PageId.DATA_PREPARATION)
    page.search.setText("needle")
    page.search.setFocus()
    before = page.table.selectionModel().selectedRows()
    QTest.keyClick(page.search, Qt.Key.Key_A, Qt.KeyboardModifier.ControlModifier)
    qapp.processEvents()
    assert page.search.selectionLength() == len("needle")
    assert page.table.selectionModel().selectedRows() == before

    invocations = []
    page.auto_triage = lambda: invocations.append("auto")
    page.export_excel = lambda: invocations.append("export")
    page.finalize = lambda: invocations.append("finalize")
    page.search.setFocus()
    QTest.keyClick(page.search, Qt.Key.Key_D, Qt.KeyboardModifier.ControlModifier)
    QTest.keyClick(page.search, Qt.Key.Key_E, Qt.KeyboardModifier.ControlModifier)
    QTest.keyClick(page.search, Qt.Key.Key_Return, Qt.KeyboardModifier.ControlModifier)
    qapp.processEvents()
    assert invocations == ["auto", "export", "finalize"]


def test_filter_excluding_all_selected_rows_leaves_zero_selection(shell, qapp, qtbot):
    shell.navigate(PageId.DATA_PREPARATION)
    page = shell.page(PageId.DATA_PREPARATION)
    selected = page.model.visible_items()[:3]
    classes = page.ctx.backend.classifications
    selected_class = classes[0]
    excluded_class = classes[1]
    for item in selected:
        shell.ctx.backend.update_item("all", item.item_id, classification=selected_class)
    page.refresh()
    _select_rows_by_click(
        page,
        qtbot,
        [
            next(
                row
                for row, item in enumerate(page.model.visible_items())
                if item.item_id == target.item_id
            )
            for target in selected
        ],
    )
    assert len(page.table.selectionModel().selectedRows()) == 3
    page.class_filter.showPopup()
    qtbot.wait(10)
    excluded_index = page.class_filter.findData(excluded_class)
    rect = page.class_filter.view().visualRect(page.class_filter.model().index(excluded_index, 0))
    QTest.mouseClick(
        page.class_filter.view().viewport(),
        Qt.MouseButton.LeftButton,
        pos=QPoint(rect.left() + 10, rect.center().y()),
    )
    qapp.processEvents()
    assert page.model.classifications == {excluded_class}
    assert page.table.selectionModel().selectedRows() == []
    assert page.selection_note.text() == "選択中の 0 件を変更"
    assert page.preview_meta.text() == "画像を選択してください"
    assert not page.usage_combo.isEnabled()


def test_usage_change_removing_selected_rows_leaves_zero_selection(shell, qapp, qtbot):
    shell.navigate(PageId.DATA_PREPARATION)
    page = shell.page(PageId.DATA_PREPARATION)
    train_rows = [
        row for row, item in enumerate(page.model.visible_items()) if item.usage == "train"
    ][:3]
    _select_rows_by_click(page, qtbot, train_rows)
    QTest.mouseClick(page.chips["val"], Qt.MouseButton.LeftButton)
    qapp.processEvents()
    assert page.table.selectionModel().selectedRows() == []
    assert page.selection_note.text() == "選択中の 0 件を変更"


def _click_menu_action(qapp, button, action):
    """ボタンからメニューを開き、指定項目の表示位置をクリックする。"""
    QTest.mouseClick(button, Qt.MouseButton.LeftButton)
    qapp.processEvents()
    menu = button.menu()
    QTest.mouseClick(menu, Qt.MouseButton.LeftButton, pos=menu.actionGeometry(action).center())
    qapp.processEvents()


def _click_table_row(qapp, page, item_id, modifiers=Qt.KeyboardModifier.NoModifier):
    """作業表の識別子セルをクリックして行を選ぶ。"""
    row = next(
        index for index, item in enumerate(page.model.visible_items()) if item.item_id == item_id
    )
    cell = page.model.index(row, 1)
    QTest.mouseClick(
        page.table.viewport(),
        Qt.MouseButton.LeftButton,
        modifiers,
        page.table.visualRect(cell).center(),
    )
    qapp.processEvents()


def _activate_mode_window(qapp, window):
    """WindowManager 管理下のモード窓へ前面化イベントを送る。"""
    QTest.qWait(550)
    qapp.sendEvent(window, QEvent(QEvent.Type.WindowDeactivate))
    qapp.sendEvent(window, QEvent(QEvent.Type.WindowActivate))
    qapp.processEvents()


def test_mask_revision_choices_are_common_and_backend_rejects_foreign_revision(shell, qapp):
    page = shell.page(PageId.DATA_PREPARATION)
    rev2 = next(item for item in page.items if "rev_002" in item.mask_revisions)
    rev1_only = next(
        item
        for item in page.items
        if item.item_id != rev2.item_id and item.mask_revisions == ["rev_001"]
    )
    _click_table_row(qapp, page, rev2.item_id)
    _click_table_row(qapp, page, rev1_only.item_id, Qt.KeyboardModifier.ControlModifier)
    assert [page.mask_combo.itemText(i) for i in range(page.mask_combo.count())] == ["rev_001"]
    with pytest.raises(ValueError, match="正解ラベル版"):
        shell.ctx.backend.update_item("all", rev1_only.item_id, selected_mask_revision="rev_002")


def test_filter_click_clears_selection_of_hidden_rows(shell, qapp):
    page = shell.page(PageId.DATA_PREPARATION)
    hidden = next(item for item in page.items if item.usage == "train")
    original_classification = hidden.classification
    _click_table_row(qapp, page, hidden.item_id)
    QTest.mouseClick(page.chips["val"], Qt.MouseButton.LeftButton)
    page.table.setFocus()
    QTest.keyClick(page.table, Qt.Key.Key_2)
    qapp.processEvents()
    assert hidden.classification == original_classification
    assert all(
        item_id in {item.item_id for item in page.model.visible_items()}
        for item_id in page._selected_ids
    )
    assert page._selected_ids == []
    assert page.selection_note.text() == "選択中の 0 件を変更"
    assert page.preview_meta.text() == "画像を選択してください"


def test_usage_and_quality_shortcuts_match_key_sequences(shell, qapp):
    page = shell.page(PageId.DATA_PREPARATION)
    item = next(item for item in page.items if item.usage == "unassigned")
    _click_table_row(qapp, page, item.item_id)
    page.table.setFocus()
    for key, usage in (
        (Qt.Key.Key_Q, "train"),
        (Qt.Key.Key_W, "val"),
        (Qt.Key.Key_E, "excluded"),
        (Qt.Key.Key_R, "unassigned"),
    ):
        QTest.keyClick(page.table, key)
        qapp.processEvents()
        assert item.usage == usage
    for key, quality in ((Qt.Key.Key_Z, "良"), (Qt.Key.Key_X, "可"), (Qt.Key.Key_C, "不良")):
        QTest.keyClick(page.table, key)
        qapp.processEvents()
        assert item.quality == quality


def test_empty_excel_export_writes_header_without_slot_exception(
    shell, qapp, tmp_path, monkeypatch
):
    pytest.importorskip("openpyxl")
    page = shell.page(PageId.DATA_PREPARATION)
    page.search.setText("no-such-item")
    output = str(tmp_path / "empty.xlsx")
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *_args: (output, "Excel"))
    _click_menu_action(qapp, page.excel_button, page.excel_button.menu().actions()[0])
    from openpyxl import load_workbook

    workbook = load_workbook(output, read_only=True)
    try:
        assert workbook.active.max_row == 1
    finally:
        workbook.close()


def test_data_menu_actions_live_as_long_as_the_page(shell):
    page = shell.page(PageId.DATA_PREPARATION)
    assert page.mask_revision_action.parent() is page
    assert page.archive_action.parent() is page
    page.refresh()


def test_error_chip_uses_all_errors_after_filtering(shell, qapp):
    page = shell.page(PageId.DATA_PREPARATION)
    expected = len(shell.ctx.backend.validate_items().errors)
    assert int(page.chips["errors"].count_label.text()) == expected
    QTest.mouseClick(page.chips["errors"], Qt.MouseButton.LeftButton)
    assert page.model.rowCount() == expected
    assert page.finalize_button.isEnabled() == (expected == 0)


def test_default_shortcuts_cover_clear_help_navigation_and_zoom(shell, qapp, monkeypatch):
    shell.navigate(PageId.DATA_PREPARATION)
    page = shell.page(PageId.DATA_PREPARATION)
    item = next(item for item in page.items if item.usage == "unassigned")
    _click_table_row(qapp, page, item.item_id)
    page.table.setFocus()
    QTest.keyClick(page.table, Qt.Key.Key_0)
    assert item.classification is None
    with patch(
        "foam_cell_analysis.gui.modes.data_preparation.page.QInputDialog.getItem",
        return_value=("分類B", True),
    ):
        QTest.keyClick(page.table, Qt.Key.Key_K)
    assert item.classification == "分類B"
    QTest.keyClick(page.table, Qt.Key.Key_F1)
    assert shell.ctx.keymap_window is not None
    assert shell.ctx.keymap_window.isVisible()
    before = page.image_view.zoom
    QTest.keyClick(page.image_view, Qt.Key.Key_Plus)
    assert page.image_view.zoom > before
    QTest.keyClick(page.image_view, Qt.Key.Key_F)
    assert page.image_view.zoom <= before
    QTest.keyClick(page.table, Qt.Key.Key_Escape)
    assert not page._selected_ids
    QTest.keyClick(page.table, Qt.Key.Key_Escape)
    assert page.model.usages is None
    QTest.keyClick(page.table, Qt.Key.Key_A, Qt.KeyboardModifier.ControlModifier)
    assert len(page.table.selectionModel().selectedRows()) == page.model.rowCount()
    all_chip_count = page.model.rowCount()
    QTest.keyClick(page.table, Qt.Key.Key_6, Qt.KeyboardModifier.AltModifier)
    qapp.processEvents()
    assert page.model.errors_only
    assert page.model.rowCount() == int(page.chips["errors"].count_label.text())
    assert all_chip_count >= page.model.rowCount()


def test_shift_space_selects_previous_unassigned(qapp, shell):
    page = shell.page(PageId.DATA_PREPARATION)
    unassigned = [item for item in page.model.visible_items() if item.usage == "unassigned"]
    assert len(unassigned) >= 2
    _click_table_row(qapp, page, unassigned[0].item_id)
    page.table.setFocus()
    QTest.keyClick(page.table, Qt.Key.Key_Space, Qt.KeyboardModifier.ShiftModifier)
    assert page._selected_ids == [unassigned[-1].item_id]
    QTest.keyClick(page.table, Qt.Key.Key_Space)
    assert page._selected_ids == [unassigned[0].item_id]


def test_imported_usage_is_visible_and_source_filter_says_all(shell, qapp):
    page = shell.page(PageId.DATA_PREPARATION)

    def finish_import(dialog):
        dialog.image_edit.setText("C:/review/images")
        dialog.mask_edit.setText("C:/review/masks")
        dialog._accept()
        return dialog.result()

    def finish_settings(dialog):
        for _folder, usage, classification, quality in dialog.rows:
            usage.setCurrentIndex(usage.findData("train"))
            classification.setCurrentIndex(classification.findData("分類A"))
            quality.setCurrentIndex(quality.findData("良"))
        dialog._accept()
        return dialog.result()

    with (
        patch.object(ImportDialog, "exec", finish_import),
        patch.object(ImportSettingsDialog, "exec", finish_settings),
    ):
        QTest.mouseClick(page.import_button, Qt.MouseButton.LeftButton)
        qapp.processEvents()
    assert page.model.rowCount() > 0
    assert page.model.usages == {"train"}
    assert page.source_combo.currentData() is None


def test_triage_returns_current_item_and_stops_at_exhaustion(shell, qapp, qtbot):
    page = shell.page(PageId.DATA_PREPARATION)
    targets = [item for item in page.items if item.usage == "unassigned"][:2]
    dialog = ContinuousTriageDialog(
        page,
        targets,
        lambda item_id, **changes: shell.ctx.backend.update_item("all", item_id, **changes),
        shell.ctx.backend,
        shell.ctx.shortcuts,
    )
    qtbot.addWidget(dialog)
    dialog.show()
    rect = dialog.filmstrip.visualItemRect(dialog.filmstrip.item(1))
    QTest.mouseClick(dialog.filmstrip.viewport(), Qt.MouseButton.LeftButton, pos=rect.center())
    assert dialog.current_item_id == targets[1].item_id
    dialog.setFocus()
    QTest.keyClick(dialog, Qt.Key.Key_Q)
    QTest.keyClick(dialog, Qt.Key.Key_Q)
    qapp.processEvents()
    assert dialog.filmstrip.count() == 0
    assert "ありません" in dialog.counter.text()
    dialog.close()


def test_excel_import_releases_source_file(shell, qapp, tmp_path, monkeypatch):
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    page = shell.page(PageId.DATA_PREPARATION)
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["id", "folder", "filename", "usage", "class", "quality", "mask"])
    item = page.items[0]
    sheet.append([item.item_id, "folder", item.source_filename, "学習", "分類A", "良", "rev_001"])
    workbook.save(path)
    workbook.close()
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *_args: (str(path), "Excel"))
    _click_menu_action(qapp, page.excel_button, page.excel_button.menu().actions()[1])
    path.rename(tmp_path / "renamed.xlsx")
    assert (tmp_path / "renamed.xlsx").exists()


def test_data_preparation_other_menu_has_required_actions(shell):
    page = shell.page(PageId.DATA_PREPARATION)
    labels = [action.text() for action in page.other_button.menu().actions()]
    assert "正解ラベル版を追加" in labels
    assert "アーカイブを作成" in labels
    assert "キー割り当て一覧" in labels


def test_multi_selection_survives_activation_and_reaches_auto_triage(shell, qapp):
    """複数選択 → ウィンドウの前面化 → 自動振り分けで、選択した全件が対象になる。"""
    from PySide6.QtWidgets import QApplication, QPushButton

    from foam_cell_analysis.gui.modes.data_preparation.dialogs import AutoTriageDialog

    shell.navigate(PageId.DATA_PREPARATION)
    page = shell.page(PageId.DATA_PREPARATION)
    window = shell.manager.window(ModeId.DATA_PREPARATION)
    window.resize(1400, 900)
    window.show()
    qapp.processEvents()
    viewport = page.table.viewport()
    first = page.table.visualRect(page.model.index(3, 1)).center()
    last = page.table.visualRect(page.model.index(10, 1)).center()
    QTest.mouseClick(viewport, Qt.MouseButton.LeftButton, pos=first)
    QTest.mouseClick(
        viewport, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.ShiftModifier, pos=last
    )
    selected = list(page._selected_ids)
    assert len(selected) == 8
    QTest.qWait(550)
    qapp.sendEvent(window, QEvent(QEvent.Type.WindowDeactivate))
    qapp.sendEvent(window, QEvent(QEvent.Type.WindowActivate))
    qapp.processEvents()
    assert page._selected_ids == selected
    assert len(page.table.selectionModel().selectedRows()) == 8
    seen = {}

    def run():
        dialog = QApplication.activeModalWidget()
        assert isinstance(dialog, AutoTriageDialog)
        seen["label"] = dialog.target_selected.text()
        QTest.mouseClick(dialog.target_selected, Qt.MouseButton.LeftButton)
        dialog.ratio.setValue(100)
        button = next(b for b in dialog.findChildren(QPushButton) if "振り分ける" in b.text())
        QTest.mouseClick(button, Qt.MouseButton.LeftButton)

    QTimer.singleShot(50, run)
    QTest.mouseClick(page.auto_button, Qt.MouseButton.LeftButton)
    qapp.processEvents()
    assert seen["label"] == "選択中（8件）"
    usages = {item.item_id: item.usage for item in shell.ctx.backend.get_working_items()}
    assert all(usages[item_id] == "val" for item_id in selected)


def test_dataset_history_row_click_opens_thumbnails_and_survives_activation(shell, qapp):
    shell.navigate(PageId.DATASET_HISTORY)
    page = shell.page(PageId.DATASET_HISTORY)
    window = shell.manager.window(ModeId.DATA_PREPARATION)
    window.show()
    qapp.processEvents()
    assert page.model.rowCount() > 0
    index = page.model.index(0, 0)
    QTest.mouseClick(
        page.table.viewport(),
        Qt.MouseButton.LeftButton,
        pos=page.table.visualRect(index).center(),
    )
    qapp.processEvents()
    selected_version = page.model.versions[0].version
    assert [row.row() for row in page.table.selectionModel().selectedRows()] == [0]
    QTest.mouseClick(page.thumbnail_button, Qt.MouseButton.LeftButton)
    qapp.processEvents()
    assert selected_version in page.thumbnail_windows
    _activate_mode_window(qapp, window)
    assert [row.row() for row in page.table.selectionModel().selectedRows()] == [0]
    assert page.table.currentIndex().row() == 0


def test_triage_close_restores_multi_selection_and_current_image(shell, qapp):
    shell.navigate(PageId.DATA_PREPARATION)
    page = shell.page(PageId.DATA_PREPARATION)
    window = shell.manager.window(ModeId.DATA_PREPARATION)
    window.show()
    qapp.processEvents()
    page.chips["all"].click()
    qapp.processEvents()
    visible = page.model.visible_items()
    chosen = [visible[row].item_id for row in (3, 5, 8)]
    for position, item_id in enumerate(chosen):
        _click_table_row(
            qapp,
            page,
            item_id,
            Qt.KeyboardModifier.NoModifier
            if position == 0
            else Qt.KeyboardModifier.ControlModifier,
        )
    _activate_mode_window(qapp, window)
    assert set(page._selected_ids) == set(chosen)
    dialog_state = {}

    def close_when_open():
        dialog = QApplication.activeModalWidget()
        if not isinstance(dialog, ContinuousTriageDialog):
            QTimer.singleShot(25, close_when_open)
            return
        dialog_state["current_id"] = dialog.current_item_id
        QTest.keyClick(dialog, Qt.Key.Key_Return)

    QTimer.singleShot(0, close_when_open)
    page.table.setFocus()
    QTest.keyClick(page.table, Qt.Key.Key_Return)
    qapp.processEvents()
    visible_after = page.model.visible_items()
    selected_after = [index.row() for index in page.table.selectionModel().selectedRows()]
    assert set(page._selected_ids) == set(chosen)
    assert "current_id" in dialog_state
    assert len(selected_after) == 3
    assert page.table.currentIndex().row() == next(
        row for row, item in enumerate(visible_after) if item.item_id == dialog_state["current_id"]
    )


def make_context(backend=None, shortcuts=None):
    """テスト用の独立したアプリケーションコンテキストを作る。"""
    from foam_cell_analysis.gui.shortcuts import ShortcutMap

    return AppContext(
        backend or MockBackend(),
        Navigator(),
        JobManager(),
        StatusBus(),
        shortcuts or ShortcutMap(mapping={}),
    )


def click_row(qapp, page, item_id):
    """対象画像の行を実際の表クリックで選択する。"""
    row = next(
        index for index, item in enumerate(page.model.visible_items()) if item.item_id == item_id
    )
    page.table.show()
    qapp.processEvents()
    rect = page.table.visualRect(page.model.index(row, 0))
    QTest.mouseClick(page.table.viewport(), Qt.MouseButton.LeftButton, pos=rect.center())
    qapp.processEvents()


def test_arrow_shortcuts_move_rows_and_ignore_search_focus(qapp):
    """左右キーで画像行を移し、検索欄では文字入力を優先する。"""
    page = DataPreparationPage(make_context())
    page.show()
    visible = page.model.visible_items()
    click_row(qapp, page, visible[1].item_id)
    QTest.keyClick(page.table, Qt.Key.Key_Left)
    assert page._selected_ids == [visible[0].item_id]
    page.search.setFocus()
    QTest.keyClick(page.search, Qt.Key.Key_Right)
    assert page._selected_ids == [visible[0].item_id]
    page.class_combo.showPopup()
    QTest.keyClick(page.class_combo, Qt.Key.Key_Left)
    assert page._selected_ids == [visible[0].item_id]
    page.class_combo.hidePopup()
    page.close()


def test_selected_auto_triage_overwrites_assigned_items_and_undoes(
    qapp, qtbot, monkeypatch, tmp_path
):
    """行クリック、Ctrl+D、選択中、実行ボタンの経路で上書きし Ctrl+Z で戻す。"""
    backend = MockBackend()
    ctx = make_context(backend)
    page = DataPreparationPage(ctx)
    page.show()
    item = next(item for item in backend.get_working_items() if item.usage == "train")
    click_row(qapp, page, item.item_id)

    def choose_and_run():
        dialog = QApplication.activeModalWidget()
        assert isinstance(dialog, AutoTriageDialog)
        QTest.mouseClick(dialog.target_selected, Qt.MouseButton.LeftButton)
        dialog.ratio.setValue(100)
        row = next(
            index
            for index in range(dialog.preview.rowCount())
            if dialog.preview.item(index, 0).text() == (item.classification or "未設定")
        )
        assert dialog.preview.item(row, 1).text().endswith("→ 0")
        assert dialog.preview.item(row, 2).text().endswith("→ 1")
        QTest.mouseClick(dialog.apply_button, Qt.MouseButton.LeftButton)

    QTimer.singleShot(0, choose_and_run)
    QTest.keyClick(page.table, Qt.Key.Key_D, Qt.KeyboardModifier.ControlModifier)
    assert item.usage == "val"
    page.undo_stack.undo()
    assert item.usage == "train"
    page.close()


def test_auto_triage_numeric_controls_and_preview_order(qapp):
    """数値欄を内容幅に保ち、分類定義順と数値セル書式を使う。"""
    backend = MockBackend()
    dialog = AutoTriageDialog(None, backend.get_working_items(), backend=backend)
    assert dialog.ratio.sizePolicy().horizontalPolicy() == QSizePolicy.Policy.Fixed
    assert dialog.seed.sizePolicy().horizontalPolicy() == QSizePolicy.Policy.Fixed
    classes = [dialog.preview.item(row, 0).text() for row in range(dialog.preview.rowCount())]
    expected = [name for name in backend.classifications if name in classes]
    if "未設定" in classes:
        expected.append("未設定")
    assert classes == expected
    if dialog.preview.rowCount():
        cell = dialog.preview.item(0, 1)
        assert cell.textAlignment() & Qt.AlignmentFlag.AlignRight
        assert cell.font().family() == numeric_font().family()
    dialog.close()


def test_single_channel_gui_and_continuous_triage_mouse_actions(qapp, qtbot):
    """取り込み画面は単一画像フォルダで、連続振り分けに用途ボタンを置く。"""
    dialog = ImportDialog()
    assert not hasattr(dialog, "channel_rows")
    assert dialog.image_edit is not None
    ctx = make_context()
    training = TrainingPage(ctx)
    assert training.fields["data.input_channels"].text() == "A（単一チャンネル）"
    assert training._collect_config()["data"]["input_channels"] == ["A"]
    training.close()
    unassigned = next(
        item for item in ctx.backend.get_working_items() if item.usage == "unassigned"
    )
    triage = ContinuousTriageDialog(
        DataPreparationPage(ctx),
        [unassigned],
        lambda item_id, **changes: ctx.backend.update_item("all", item_id, **changes),
        ctx.backend,
        ctx.shortcuts,
        ctx.display,
    )
    triage.show()
    qapp.processEvents()
    assert set(triage.usage_buttons) == {"train", "val", "excluded", "unassigned"}
    initial_usage = unassigned.usage
    QTest.mouseClick(triage.usage_buttons["train"][0], Qt.MouseButton.LeftButton)
    assert unassigned.usage != initial_usage
    QTest.mouseClick(triage.target_filtered, Qt.MouseButton.LeftButton)
    QTest.keyClick(triage, Qt.Key.Key_R)
    assert (
        next(
            item for item in ctx.backend.get_working_items() if item.item_id == unassigned.item_id
        ).usage
        == "unassigned"
    )
    triage.close()


def test_dataset_history_opens_read_only_thumbnail_window(qapp):
    """版履歴の選択とボタン操作で、版の全画像を非モーダル表示する。"""
    ctx = make_context()
    page = DatasetHistoryPage(ctx)
    page.show()
    page.table.selectRow(0)
    qapp.processEvents()
    QTest.mouseClick(page.thumbnail_button, Qt.MouseButton.LeftButton)
    version = page.model.versions[0]
    window = page.thumbnail_windows[version.version]
    assert not window.isModal()
    model = window.models[version.version]
    assert model.rowCount() == version.n_images
    if version.purpose == "train" and version.base_validation_version:
        assert window.tabs.count() == 2
        assert version.base_validation_version in window.tabs.tabText(1)
    index = page.model.index(0, 0)
    QTest.mouseDClick(
        page.table.viewport(),
        Qt.MouseButton.LeftButton,
        pos=page.table.visualRect(index).center(),
    )
    assert page.thumbnail_windows[version.version] is window
    window.close()
    page.close()


def test_finalize_reason_button_filters_error_rows(qapp):
    """確定不可の理由ボタンからエラー行だけに絞り込める。"""
    backend = MockBackend()
    item = next(item for item in backend.get_working_items() if item.usage == "train")
    backend.update_item("all", item.item_id, classification=None)
    page = DataPreparationPage(make_context(backend))
    page.show()
    assert not page.finalize_error_button.isHidden()
    QTest.mouseClick(page.finalize_error_button, Qt.MouseButton.LeftButton)
    assert page.model.errors_only
    assert all(
        item.item_id in {error.item_id for error in backend.validate_items().errors}
        for item in page.model.visible_items()
    )
    dialog = DatasetFinalizeDialog(None, backend)
    assert not dialog.error_summary_button.isHidden()
    dialog.close()
    page.close()


def test_display_toggle_ignores_arrow_keys_in_continuous_triage(qapp):
    """連続振り分けで ← → は画像移動だけに効き、表示切り替えは変わらない。"""
    page = DataPreparationPage(make_context())
    dialog = ContinuousTriageDialog(
        page,
        page.model.visible_items(),
        lambda *_args, **_kwargs: None,
        page.ctx.backend,
        page.ctx.shortcuts,
    )
    dialog.show()
    qapp.processEvents()
    toggle = dialog.display_toggle
    before = toggle.is_alternate
    QTest.mouseClick(toggle.raw_button, Qt.MouseButton.LeftButton)
    assert not toggle.raw_button.hasFocus()
    QTest.keyClick(dialog, Qt.Key.Key_Right)
    QTest.keyClick(dialog, Qt.Key.Key_Left)
    assert toggle.is_alternate is False
    assert toggle.raw_button.focusPolicy() == Qt.FocusPolicy.NoFocus
    assert before is True
    dialog.close()
    page.close()
